"""Unit tests for eoa.report.trends (FR-4.3 trend detection, synthetic rows, no DB) and the
weekly/monthly report builders (eoa.report.weekly / eoa.report.monthly), with every DB- and
LLM-touching collector monkeypatched so these run with no Postgres and no Ollama.

Run with: ``PYTHONPATH=agent python -m pytest tests/unit/test_report_weekly_monthly.py -q``
"""

from __future__ import annotations

import datetime as dt

import docx
import pytest

from eoa.llm.schemas.analysis import OutlookIndicator, Sentence, StructuredSection
from eoa.llm.schemas.reports import MonthlyReportDraft, WeeklyReportDraft, WeeklyTrendSection
from eoa.report import monthly, trends, weekly

# --------------------------------------------------------------------------
# trends.py: pure-function trend detectors, synthetic rows, no DB
# --------------------------------------------------------------------------


def test_entity_clusters_from_rows_builds_a_trend():
    rows = [{"entity": "Elbit Systems", "domain": "airborne_pods", "item_ids": [3, 1, 2], "n": 3}]
    out = trends._entity_clusters_from_rows(rows)
    assert len(out) == 1
    t = out[0]
    assert t["kind"] == "entity_cluster"
    assert t["entities"] == ["Elbit Systems"]
    assert t["evidence_item_ids"] == [1, 2, 3]
    assert 1 <= t["strength"] <= 5


def test_entity_clusters_filters_below_threshold():
    rows = [{"entity": "X", "domain": "d", "item_ids": [1, 2], "n": 2}]
    assert trends._entity_clusters_from_rows(rows) == []


def test_domain_surge_detected_when_ratio_exceeds_2x_baseline():
    counts = {"c_uas": 10}
    baseline = {"c_uas": 2.0}
    items_by_domain = {"c_uas": [1, 2, 3, 4, 5]}
    out = trends._domain_surges_from_counts(counts, baseline, items_by_domain)
    assert len(out) == 1
    assert out[0]["kind"] == "domain_surge"
    assert out[0]["strength"] == 5  # ratio 5x -> top strength


def test_domain_surge_falls_back_to_floor_when_no_baseline_history():
    counts = {"computer_vision": 3}
    baseline: dict[str, float] = {}
    items_by_domain = {"computer_vision": [10, 11, 12]}
    out = trends._domain_surges_from_counts(counts, baseline, items_by_domain)
    assert len(out) == 1


def test_domain_surge_not_triggered_below_2x_ratio():
    counts = {"c_uas": 4}
    baseline = {"c_uas": 3.0}  # ratio 1.33, not >= 2x
    out = trends._domain_surges_from_counts(counts, baseline, {"c_uas": [1, 2, 3, 4]})
    assert out == []


def test_market_convergence_from_rows_builds_a_trend():
    rows = [{"subdomain": "targeting_pods", "n": 2, "item_ids": [2, 1], "parties": ["Elbit", "Rafael"]}]
    out = trends._convergence_from_rows(rows)
    assert len(out) == 1
    assert out[0]["kind"] == "market_convergence"
    assert out[0]["entities"] == ["Elbit", "Rafael"]
    assert out[0]["evidence_item_ids"] == [1, 2]


def test_market_convergence_filters_below_threshold():
    rows = [{"subdomain": "x", "n": 1, "item_ids": [1], "parties": ["A"]}]
    assert trends._convergence_from_rows(rows) == []


def test_tech_race_from_rows_builds_a_trend():
    rows = [{"subdomain": "c_uas_effectors", "item_ids": [1, 2], "companies": ["Rafael", "Elbit"]}]
    out = trends._tech_race_from_rows(rows)
    assert len(out) == 1
    assert out[0]["kind"] == "tech_race"
    assert set(out[0]["entities"]) == {"Rafael", "Elbit"}


def test_tech_race_filters_single_company():
    rows = [{"subdomain": "x", "item_ids": [1], "companies": ["OnlyOne"]}]
    assert trends._tech_race_from_rows(rows) == []


def test_detect_trends_combines_all_four_kinds(monkeypatch):
    monkeypatch.setattr(
        trends,
        "_entity_cluster_rows",
        lambda s, e: [{"entity": "A", "domain": "d1", "item_ids": [1, 2, 3], "n": 3}],
    )
    monkeypatch.setattr(trends, "_domain_counts", lambda s, e: {"d2": 10})
    monkeypatch.setattr(trends, "_domain_baseline_counts", lambda s, e: {"d2": 1.0})
    monkeypatch.setattr(trends, "_domain_item_ids", lambda s, e: {"d2": [4, 5, 6]})
    monkeypatch.setattr(
        trends,
        "_convergence_rows",
        lambda s, e: [{"subdomain": "sd1", "n": 2, "item_ids": [7, 8], "parties": ["X", "Y"]}],
    )
    monkeypatch.setattr(
        trends,
        "_tech_race_rows",
        lambda s, e: [{"subdomain": "sd2", "item_ids": [9, 10], "companies": ["P", "Q"]}],
    )
    out = trends.detect_trends((dt.date(2026, 8, 25), dt.date(2026, 8, 31)))
    kinds = {t["kind"] for t in out}
    assert kinds == {"entity_cluster", "domain_surge", "market_convergence", "tech_race"}
    assert [t["strength"] for t in out] == sorted((t["strength"] for t in out), reverse=True)


def test_weekly_stats_assembles_all_four_aggregates(monkeypatch):
    monkeypatch.setattr(
        trends, "_items_by_domain_level", lambda s, e: [{"domain": "d", "level": "red", "n": 1}]
    )
    monkeypatch.setattr(
        trends,
        "_top_entities_with_delta",
        lambda s, e, limit=10: [{"entity": "A", "mentions": 3, "delta": 1}],
    )
    monkeypatch.setattr(trends, "_events_by_kind", lambda s, e: [{"kind": "contract_award", "n": 2}])
    monkeypatch.setattr(trends, "_deep_search_outcomes", lambda s, e: [{"outcome": "found", "n": 1}])
    out = trends.weekly_stats((dt.date(2026, 8, 25), dt.date(2026, 8, 31)))
    assert set(out) == {"items_by_domain_level", "top_entities", "events_by_kind", "deep_search_outcomes"}
    assert out["top_entities"][0]["entity"] == "A"


# --------------------------------------------------------------------------
# weekly.py: build_weekly renders docx with the trend section + calendar table, QA passes
# --------------------------------------------------------------------------

WEEKLY_ITEMS = [
    {
        "id": 101,
        "n": 1,
        "title": "Elbit wins targeting pod contract",
        "domain": "airborne_pods",
        "source_name": "Defense News",
        "url": "https://example.com/1",
        "published_at": dt.date(2026, 8, 30),
        "level": "red",
        "summary_he": "אלביט מערכות זכתה בחוזה של 50 מיליון דולר [1].",
        "so_what_he": "צעד משמעותי בתחרות.",
    },
    {
        "id": 102,
        "n": 2,
        "title": "Rafael launches new C-UAS system",
        "domain": "c_uas",
        "source_name": "Janes",
        "url": "https://example.com/2",
        "published_at": dt.date(2026, 8, 31),
        "level": "orange",
        "summary_he": "רפאל השיקה מערכת נגד כטבמים חדשה [2].",
        "so_what_he": "תחרות מוגברת.",
    },
]


def _weekly_draft_fixture() -> WeeklyReportDraft:
    return WeeklyReportDraft(
        exec_summary=[
            Sentence(text_he="אלביט מערכות זכתה בחוזה של 50 מיליון דולר.", cites=[1]),
            Sentence(text_he="רפאל השיקה מערכת נגד כטבמים חדשה.", cites=[2]),
        ],
        trends=[
            WeeklyTrendSection(
                title_he="מגמה: פעילות מוגברת סביב Elbit Systems",
                sentences=[
                    Sentence(
                        text_he="נרשמה עלייה בפעילות סביב אלביט מערכות בתחום הפודים האוויריים.", cites=[1]
                    )
                ],
            ),
        ],
        sections=[
            StructuredSection(
                title_he="פודים ומטענים אוויריים",
                domain="airborne_pods",
                sentences=[Sentence(text_he="אלביט מערכות זכתה בחוזה בהיקף 50 מיליון דולר.", cites=[1])],
            ),
        ],
        outlook=[
            OutlookIndicator(text_he="להערכתנו מגמת ההשקות תימשך ברבעון הקרוב.", cites=[], is_assessment=True)
        ],
        open_points_he=["האם ידועות תוכניות המשך?"],
    )


@pytest.fixture
def patch_weekly_collectors(monkeypatch, tmp_path):
    monkeypatch.setattr(weekly, "collect_week_items", lambda s, e: [dict(it) for it in WEEKLY_ITEMS])
    monkeypatch.setattr(
        weekly, "collect_yellow_domain_summary", lambda s, e: [{"domain": "computer_vision", "n": 4}]
    )
    monkeypatch.setattr(weekly, "collect_events", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_deep_search", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_open_clarifications", lambda: [])
    monkeypatch.setattr(
        weekly.trends_mod,
        "detect_trends",
        lambda period: [
            {
                "kind": "entity_cluster",
                "title_he": "מגמה: פעילות מוגברת סביב Elbit Systems",
                "evidence_item_ids": [101],
                "entities": ["Elbit Systems"],
                "strength": 4,
            }
        ],
    )
    monkeypatch.setattr(
        weekly,
        "collect_meta_summary",
        lambda s, e: {
            "lessons": [{"id": 1, "text": "פריטי RFI סווגו כ-yellow באופן שיטתי מדי", "created_at": None}],
            "feedback_total": 2,
            "feedback_deltas": [],
        },
    )
    monkeypatch.setattr(
        weekly,
        "upcoming_conferences",
        lambda days=90: [
            {
                "name": "DSEI",
                "start_date": dt.date(2026, 9, 10),
                "end_date": dt.date(2026, 9, 13),
                "city": "London",
                "relevance": 5,
            }
        ],
    )
    monkeypatch.setattr(weekly, "draft_weekly", lambda *a, **k: _weekly_draft_fixture())
    monkeypatch.setattr(weekly, "_persist_report", lambda *a, **k: 999)
    monkeypatch.setattr(weekly, "_report_path", lambda period_end, ext: tmp_path / f"weekly.{ext}")
    return tmp_path


def test_build_weekly_qa_passes_on_fixture_draft(patch_weekly_collectors):
    paths = weekly.build_weekly(period_end=dt.date(2026, 9, 4))
    assert paths.qa.passed
    assert paths.qa.errors == []
    assert paths.report_id == 999


def test_build_weekly_renders_docx_with_trend_section_and_calendar_table(patch_weekly_collectors):
    paths = weekly.build_weekly(period_end=dt.date(2026, 9, 4))
    assert paths.docx.exists()
    doc = docx.Document(str(paths.docx))
    heading_texts = {p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"}
    assert "לוח 90 הימים הקרובים" in heading_texts
    assert "מגמה: פעילות מוגברת סביב Elbit Systems" in heading_texts
    assert "סיכום מטא שבועי — משוב משתמש (FR-11.4)" in heading_texts

    calendar_table = next(
        (t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["שם", "תאריכים", "עיר", "רלוונטיות"]),
        None,
    )
    assert calendar_table is not None
    assert calendar_table.rows[1].cells[0].text == "DSEI"


def test_build_weekly_suppresses_empty_meta_summary_section(patch_weekly_collectors, monkeypatch):
    """U13: a heading over a "nothing happened" placeholder line is worse than no section at all —
    when there are no lessons, no feedback deltas, and zero feedback total, the meta-summary
    section must not be rendered."""
    monkeypatch.setattr(
        weekly,
        "collect_meta_summary",
        lambda s, e: {"lessons": [], "feedback_total": 0, "feedback_deltas": []},
    )
    paths = weekly.build_weekly(period_end=dt.date(2026, 9, 4))
    doc = docx.Document(str(paths.docx))
    heading_texts = {p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"}
    assert "סיכום מטא שבועי — משוב משתמש (FR-11.4)" not in heading_texts


def test_build_weekly_md_and_html_contain_trend_and_calendar(patch_weekly_collectors):
    paths = weekly.build_weekly(period_end=dt.date(2026, 9, 4))
    md_text = paths.md.read_text(encoding="utf-8")
    html_text = paths.html.read_text(encoding="utf-8")
    assert "לוח 90 הימים הקרובים" in md_text
    assert "DSEI" in md_text
    # the heading text survives, but the embedded "90" is wrapped in its own <bdi dir="ltr"> span
    # (F10: html bidi handling for a Latin/digit run inside RTL Hebrew text), so check the Hebrew
    # part either side of the number rather than the whole literal string as one run.
    assert "לוח " in html_text
    assert "הימים הקרובים" in html_text
    assert "DSEI" in html_text


def test_build_weekly_no_items_skips_llm_and_still_persists(monkeypatch, tmp_path):
    monkeypatch.setattr(weekly, "collect_week_items", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_yellow_domain_summary", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_events", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_deep_search", lambda s, e, limit=None: [])
    monkeypatch.setattr(weekly, "collect_open_clarifications", lambda: [])
    monkeypatch.setattr(weekly.trends_mod, "detect_trends", lambda period: [])
    monkeypatch.setattr(
        weekly,
        "collect_meta_summary",
        lambda s, e: {"lessons": [], "feedback_total": 0, "feedback_deltas": []},
    )
    monkeypatch.setattr(weekly, "upcoming_conferences", lambda days=90: [])
    monkeypatch.setattr(weekly, "_persist_report", lambda *a, **k: 1)
    monkeypatch.setattr(weekly, "_report_path", lambda period_end, ext: tmp_path / f"weekly.{ext}")

    def _fail_if_called(*a, **k):
        raise AssertionError("chat_structured should not be called with zero items")

    monkeypatch.setattr("eoa.report.weekly.chat_structured", _fail_if_called)

    paths = weekly.build_weekly(period_end=dt.date(2026, 9, 4))
    assert paths.qa.passed


# --------------------------------------------------------------------------
# monthly.py: build_monthly renders a players-map table per domain
# --------------------------------------------------------------------------

MONTH_ITEMS = [
    {
        "id": 201,
        "n": 1,
        "title": "IAI wins naval radar deal",
        "domain": "naval_surveillance",
        "source_name": "Naval News",
        "url": "https://example.com/201",
        "published_at": dt.date(2026, 8, 5),
        "level": "red",
        "summary_he": "IAI זכתה בחוזה חדש לאספקת מערכת EO ימית [1].",
        "so_what_he": "מחזק את מעמדה התחרותי.",
    },
]


def _monthly_draft_fixture() -> MonthlyReportDraft:
    return MonthlyReportDraft(
        exec_summary=[Sentence(text_he="IAI זכתה בחוזה חדש לאספקת מערכת EO ימית.", cites=[1])],
        trends=[],
        sections=[
            StructuredSection(
                title_he="תצפית ימית",
                domain="naval_surveillance",
                sentences=[Sentence(text_he="IAI זכתה בחוזה חדש.", cites=[1])],
            ),
        ],
        outlook=[OutlookIndicator(text_he="להערכתנו המגמה תימשך.", cites=[], is_assessment=True)],
        open_points_he=[],
    )


@pytest.fixture
def patch_monthly_collectors(monkeypatch, tmp_path):
    monkeypatch.setattr(monthly, "collect_month_items", lambda s, e: [dict(it) for it in MONTH_ITEMS])
    monkeypatch.setattr(monthly, "collect_yellow_domain_summary", lambda s, e, limit=None: [])
    monkeypatch.setattr(monthly, "collect_events", lambda s, e, limit=None: [])
    monkeypatch.setattr(monthly, "collect_deep_search", lambda s, e, limit=None: [])
    monkeypatch.setattr(monthly, "collect_open_clarifications", lambda: [])
    monkeypatch.setattr(monthly.trends_mod, "detect_trends", lambda period: [])
    monkeypatch.setattr(monthly, "collect_previous_monthly_trends", lambda period_start: [])
    monkeypatch.setattr(monthly, "draft_monthly", lambda *a, **k: _monthly_draft_fixture())
    monkeypatch.setattr(
        monthly,
        "players_map",
        lambda: {
            "naval_surveillance": [
                {"entity_id": 1, "name": "IAI", "COMPETITOR_OF": 2, "SUPPLIER_OF": 1, "PARTNER_OF": 3}
            ]
        },
    )
    monkeypatch.setattr(monthly, "top_events_by_amount", lambda s, e, limit=10: [])
    monkeypatch.setattr(monthly, "full_horizon_table", lambda: [])
    monkeypatch.setattr(monthly, "watchlist_changes", lambda s, e, limit=None: [])
    monkeypatch.setattr(monthly, "_persist_report", lambda *a, **k: 888)
    monkeypatch.setattr(monthly, "_report_path", lambda period_end, ext: tmp_path / f"monthly.{ext}")
    return tmp_path


def test_build_monthly_qa_passes_and_persists(patch_monthly_collectors):
    paths = monthly.build_monthly(period_end=dt.date(2026, 8, 31))
    assert paths.qa.passed
    assert paths.report_id == 888


def test_build_monthly_renders_players_map_table_per_domain(patch_monthly_collectors):
    paths = monthly.build_monthly(period_end=dt.date(2026, 8, 31))
    doc = docx.Document(str(paths.docx))
    heading_texts = [p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    assert any(t.startswith("נוף תחרותי") for t in heading_texts)

    players_table = next(
        (t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["ישות", "מתחרים", "ספקים", "שותפים"]),
        None,
    )
    assert players_table is not None
    row = [c.text for c in players_table.rows[1].cells]
    assert row == ["IAI", "2", "1", "3"]


def test_month_range_defaults_to_previous_calendar_month(monkeypatch):
    monkeypatch.setattr(monthly, "_today_jerusalem", lambda: dt.date(2026, 9, 4))
    start, end = monthly._month_range(None)
    assert start == dt.date(2026, 8, 1)
    assert end == dt.date(2026, 8, 31)


def test_month_range_explicit_end_computes_full_month_bounds():
    start, end = monthly._month_range(dt.date(2026, 2, 15))
    assert start == dt.date(2026, 2, 1)
    assert end == dt.date(2026, 2, 28)
