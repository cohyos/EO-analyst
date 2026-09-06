"""Round 5 QA-loop repair (P5, patent survey benchmark gaps) -- unit tests for the six items in
docs/REPORT_TEMPLATE_BENCHMARK.md 2.5/3.5/4 (items 8-9) and docs/qa/loop/round_3_fixes.md's patent
nits / round_3_judge.md's D8 #4 finding:

1. deterministic "שיטה והיקף" (methodology & scope) box before the executive summary, with the
   coverage caveat removed from `exec_summary` itself.
2. bogus (region/country/generic-noun) assignees dropped, never profiled.
3. TF-IDF-lite sub-clustering of the "לא מסווג" bucket when CPC is missing.
4. `_verify_relationship_edges` actually drops an untraceable edge (the Anduril-Elbit "Sigma 155"
   case) and dedupes a canonical-party-pair duplicate.
5. `priority` + `confidence` on every business implication (schema/render).
6. CPC x assignee white-space matrix table, omitted honestly when either axis is empty.

Mirrors tests/unit/test_patents_survey.py's/test_patents_round3.py's own "stub every DB/network
call" convention -- no live DB/Ollama/HTTP calls anywhere in this file.
"""

from __future__ import annotations

import datetime as dt

import docx
import pytest
from pydantic import ValidationError

from eoa.llm.schemas.patents import PatentBizAction
from eoa.patents.cluster import (
    UNCLASSIFIED_KEY,
    UNCLASSIFIED_LABEL_HE,
    UNCLASSIFIED_MAX_SUBCLUSTERS,
    cluster_patents,
    tfidf_subcluster_unclassified,
)
from eoa.patents.render import (
    insert_section_before_html_summary,
    insert_section_before_md_summary,
    insert_section_before_summary_docx,
)
from eoa.patents.survey import (
    METHODOLOGY_BOX_TITLE_HE,
    _cpc_assignee_matrix,
    _cpc_coverage,
    _date_range_he,
    _is_real_company_assignee,
    _name_in_text,
    _priority_confidence_prefix_he,
    _sources_scanned_he,
    _verify_relationship_edges,
    business_implications_table,
    methodology_box_lines_he,
)

# --------------------------------------------------------------------------
# item 1: "שיטה והיקף" methodology/scope box
# --------------------------------------------------------------------------


class TestCpcCoverage:
    def test_partial_coverage(self):
        rows = [{"cpc": ["G01"]}, {"cpc": []}, {"cpc": ["H01"]}]
        missing, total, coverage = _cpc_coverage(rows)
        assert (missing, total) == (1, 3)
        assert coverage == pytest.approx(2 / 3)

    def test_empty_sample_is_vacuously_full_coverage(self):
        assert _cpc_coverage([]) == (0, 0, 1.0)

    def test_full_coverage(self):
        rows = [{"cpc": ["G01"]}, {"cpc": ["H01"]}]
        assert _cpc_coverage(rows) == (0, 2, 1.0)


class TestDateRangeHe:
    def test_uses_publication_date_when_present(self):
        rows = [
            {"publication_date": dt.date(2020, 1, 1)},
            {"publication_date": dt.date(2022, 6, 1)},
        ]
        assert _date_range_he(rows) == (dt.date(2020, 1, 1), dt.date(2022, 6, 1))

    def test_falls_back_to_filing_then_priority_date(self):
        rows = [
            {"publication_date": None, "filing_date": dt.date(2019, 3, 1), "priority_date": None},
            {"publication_date": None, "filing_date": None, "priority_date": dt.date(2021, 5, 5)},
        ]
        assert _date_range_he(rows) == (dt.date(2019, 3, 1), dt.date(2021, 5, 5))

    def test_none_when_no_dates_at_all(self):
        rows = [{"publication_date": None, "filing_date": None, "priority_date": None}]
        assert _date_range_he(rows) is None

    def test_none_for_empty_rows(self):
        assert _date_range_he([]) is None


class TestSourcesScannedHe:
    def test_keyless_label_when_no_structured_keys_configured(self, monkeypatch):
        monkeypatch.setattr("eoa.patents.survey.scan_mod.structured_sources_configured", lambda: False)
        text = _sources_scanned_he()
        assert "Google Patents" in text

    def test_structured_label_when_keys_configured(self, monkeypatch):
        monkeypatch.setattr("eoa.patents.survey.scan_mod.structured_sources_configured", lambda: True)
        text = _sources_scanned_he()
        assert "EPO OPS" in text and "PatentsView" in text


class TestMethodologyBoxLinesHe:
    def _lines(self, **overrides):
        base = dict(
            topic="thermal imaging seeker",
            sources_scanned_he="Google Patents (חיפוש חסר-מפתחות)",
            date_range=(dt.date(2018, 1, 1), dt.date(2024, 1, 1)),
            n_fresh=12,
            n_stored=5,
            assignee_with=3,
            assignee_total=17,
            cpc_with=0,
            cpc_total=17,
            caveat_he=None,
        )
        base.update(overrides)
        return methodology_box_lines_he(**base)

    def test_carries_query_and_record_counts(self):
        lines = self._lines()
        joined = " | ".join(lines)
        assert "thermal imaging seeker" in joined
        assert "12" in joined and "5" in joined  # fresh vs stored
        assert "17" in joined  # total records

    def test_assignee_coverage_prominent_tag_shape(self):
        lines = self._lines(assignee_with=3, assignee_total=17)
        assert any(
            line.startswith("כיסוי נתוני מקצה: ") and "18%" in line and "(3/17)" in line for line in lines
        )

    def test_cpc_coverage_line_present(self):
        lines = self._lines(cpc_with=0, cpc_total=17)
        assert any(line.startswith("כיסוי קודי CPC: ") and "0%" in line for line in lines)

    def test_caveat_appended_as_last_line_when_given(self):
        caveat = "ל-14 מתוך 17 הפטנטים אין נתוני מקצה — לא ניתן להסיק בלעדיות או נתח שוק."
        lines = self._lines(caveat_he=caveat)
        assert lines[-1] == caveat

    def test_no_caveat_line_when_none(self):
        lines = self._lines(caveat_he=None)
        assert not any("בלעדיות" in line for line in lines)

    def test_missing_date_range_is_disclosed_not_fabricated(self):
        lines = self._lines(date_range=None)
        assert any("לא זמין" in line for line in lines)

    def test_empty_sample_shows_100_percent_coverage_vacuously(self):
        lines = self._lines(assignee_with=0, assignee_total=0, cpc_with=0, cpc_total=0)
        assert any("100%" in line and "(0/0)" in line for line in lines if "מקצה" in line)


class TestInsertMethodologyBoxMd:
    def test_inserts_before_summary_heading_as_bullets(self):
        text = "# כותרת\n\n## תקציר מנהלים\n\nתוכן.\n\n## נספח מקורות\n\nרשומות\n"
        out = insert_section_before_md_summary(text, METHODOLOGY_BOX_TITLE_HE, ["שורה א", "שורה ב"])
        lines = out.split("\n")
        assert lines.index(f"## {METHODOLOGY_BOX_TITLE_HE}") < lines.index("## תקציר מנהלים")
        assert "- שורה א" in out and "- שורה ב" in out

    def test_unchanged_when_no_summary_heading_found(self):
        text = "# כותרת\n\nללא תקציר.\n"
        out = insert_section_before_md_summary(text, METHODOLOGY_BOX_TITLE_HE, ["שורה"])
        assert out == text


class TestInsertMethodologyBoxHtml:
    def test_inserts_before_summary_h2_as_list(self):
        text = '<h1>כותרת</h1>\n<h2 id="s0">תקציר מנהלים</h2>\n<p>תוכן</p>\n'
        out = insert_section_before_html_summary(text, METHODOLOGY_BOX_TITLE_HE, ["שורה א", "שורה ב"])
        lines = out.split("\n")
        idx_box = next(i for i, line in enumerate(lines) if METHODOLOGY_BOX_TITLE_HE in line)
        idx_summary = next(i for i, line in enumerate(lines) if "תקציר מנהלים" in line)
        assert idx_box < idx_summary
        assert "שורה א" in out and "שורה ב" in out

    def test_unchanged_when_no_summary_heading_found(self):
        text = "<h1>כותרת</h1>\n<p>ללא תקציר</p>\n"
        out = insert_section_before_html_summary(text, METHODOLOGY_BOX_TITLE_HE, ["שורה"])
        assert out == text


class TestInsertMethodologyBoxDocx:
    def test_inserts_heading1_and_bullets_before_exec_summary(self):
        doc = docx.Document()
        doc.add_paragraph("כותרת", style="Title")
        doc.add_paragraph("תקציר מנהלים", style="Heading 1")
        doc.add_paragraph("תוכן התקציר.")
        insert_section_before_summary_docx(doc, METHODOLOGY_BOX_TITLE_HE, ["שורה א", "שורה ב"])
        texts = [p.text for p in doc.paragraphs]
        idx_box = texts.index(METHODOLOGY_BOX_TITLE_HE)
        idx_summary = texts.index("תקציר מנהלים")
        assert idx_box < idx_summary
        assert texts[idx_box + 1] == "שורה א"
        assert texts[idx_box + 2] == "שורה ב"
        assert idx_box + 2 < idx_summary

    def test_noop_when_no_exec_summary_heading(self):
        doc = docx.Document()
        doc.add_paragraph("כותרת", style="Title")
        before = [p.text for p in doc.paragraphs]
        insert_section_before_summary_docx(doc, METHODOLOGY_BOX_TITLE_HE, ["שורה"])
        after = [p.text for p in doc.paragraphs]
        assert before == after


# --------------------------------------------------------------------------
# item 2: bogus assignees (region/country/generic-noun) dropped
# --------------------------------------------------------------------------


class TestGenericAndCountryAssigneesRejected:
    @pytest.mark.parametrize("name", ["Inc", "Inc.", "Ltd", "Ltd.", "LLC", "Corp", "Corporation", "Systems"])
    def test_bare_generic_corporate_suffix_rejected(self, name):
        assert _is_real_company_assignee(name) is False

    @pytest.mark.parametrize("name", ["United States", 'ארה"ב', "USA"])
    def test_plain_country_name_rejected(self, name):
        assert _is_real_company_assignee(name) is False

    def test_real_company_with_generic_looking_but_distinct_name_still_counts(self):
        # A real company name is never *exactly* one of the bare generic tokens on its own --
        # "Systems Engineering Solutions Ltd" (a whole, distinct string) must not be caught by the
        # bare-token regex, which only matches when the *entire* trimmed string is generic.
        assert _is_real_company_assignee("Systems Engineering Solutions Ltd") is True


# --------------------------------------------------------------------------
# item 3: TF-IDF-lite sub-clustering of the unclassified bucket
# --------------------------------------------------------------------------


class TestTfidfSubclusterUnclassified:
    def test_empty_rows_returns_empty(self):
        assert tfidf_subcluster_unclassified([]) == []

    def test_distinct_topics_land_in_different_subclusters(self):
        rows = [
            {
                "n": 1,
                "id": 1,
                "title": "thermal infrared sensor array",
                "abstract": "infrared detector array pixel",
            },
            {
                "n": 2,
                "id": 2,
                "title": "thermal infrared detector pixel",
                "abstract": "infrared sensor array",
            },
            {
                "n": 3,
                "id": 3,
                "title": "battery charging circuit",
                "abstract": "power management circuit charging",
            },
            {
                "n": 4,
                "id": 4,
                "title": "battery charging power circuit",
                "abstract": "charging circuit management",
            },
        ]
        clusters = tfidf_subcluster_unclassified(rows, similarity_threshold=0.2)
        assert len(clusters) >= 2
        all_ns = sorted(n for c in clusters for n in c["patent_ns"])
        assert all_ns == [1, 2, 3, 4]
        # the two thermal-imaging rows must not split from each other into singleton clusters
        thermal_cluster = next(c for c in clusters if 1 in c["patent_ns"])
        assert 2 in thermal_cluster["patent_ns"]

    def test_never_exceeds_max_clusters(self):
        rows = [
            {"n": i, "id": i, "title": f"unique topic {i} alpha beta gamma", "abstract": ""}
            for i in range(20)
        ]
        clusters = tfidf_subcluster_unclassified(rows, max_clusters=3)
        assert len(clusters) <= 3
        all_ns = sorted(n for c in clusters for n in c["patent_ns"])
        assert all_ns == list(range(20))

    def test_row_with_no_tokens_still_lands_somewhere(self):
        rows = [{"n": 1, "id": 1, "title": "", "abstract": ""}]
        clusters = tfidf_subcluster_unclassified(rows)
        assert sum(len(c["patent_ns"]) for c in clusters) == 1
        assert clusters[0]["label_terms"] == []

    def test_default_max_clusters_matches_module_constant(self):
        rows = [
            {"n": i, "id": i, "title": f"totally distinct subject matter number {i}", "abstract": ""}
            for i in range(50)
        ]
        clusters = tfidf_subcluster_unclassified(rows)
        assert len(clusters) <= UNCLASSIFIED_MAX_SUBCLUSTERS


class TestClusterPatentsUnclassifiedSubclustering:
    def test_multiple_subclusters_get_distinct_labelled_keys(self):
        rows = [
            {
                "n": 1,
                "id": 1,
                "cpc": [],
                "assignees": [],
                "title": "thermal infrared sensor array pixel",
                "abstract": "infrared detector array pixel readout",
                "publication_date": None,
            },
            {
                "n": 2,
                "id": 2,
                "cpc": [],
                "assignees": [],
                "title": "thermal infrared detector pixel array",
                "abstract": "infrared sensor readout pixel array",
                "publication_date": None,
            },
            {
                "n": 3,
                "id": 3,
                "cpc": [],
                "assignees": [],
                "title": "battery charging power circuit",
                "abstract": "charging circuit management power",
                "publication_date": None,
            },
            {
                "n": 4,
                "id": 4,
                "cpc": [],
                "assignees": [],
                "title": "battery power charging circuit management",
                "abstract": "power circuit charging management",
                "publication_date": None,
            },
        ]
        clusters = cluster_patents(rows)
        unclassified = [c for c in clusters if c.key.startswith(UNCLASSIFIED_KEY)]
        assert len(unclassified) >= 2
        assert all(c.label_he != UNCLASSIFIED_LABEL_HE for c in unclassified)
        assert all(c.label_he.startswith("אשכול נושאי:") for c in unclassified)  # round 6: term-derived label, never "לא מסווג"
        total = sum(c.total for c in unclassified)
        assert total == 4

    def test_single_subcluster_keeps_flat_label(self):
        rows = [
            {
                "n": 1,
                "id": 1,
                "cpc": [],
                "assignees": [],
                "title": "",
                "abstract": "",
                "publication_date": None,
            }
        ]
        clusters = cluster_patents(rows)
        assert clusters[0].label_he == UNCLASSIFIED_LABEL_HE


# --------------------------------------------------------------------------
# item 4: relationship-edge verification -- Anduril/Elbit "Sigma 155" case
# --------------------------------------------------------------------------


class TestVerifyRelationshipEdgesRound5:
    def test_regression_stray_word_no_longer_falsely_names_anduril(self):
        """The exact round-3-judge regression (round_3_judge.md D8 #4): a source article about an
        unrelated Elbit personnel appointment merely contains the ordinary English word "anvil" --
        which happens to be one of Anduril's own *strict* watchlist aliases (config/watchlist.yaml)
        -- and the round-3 substring check wrongly read that as "the source names Anduril",
        letting the untraceable Anduril<->Elbit "Sigma 155" edge survive verification. It must now
        be dropped: Anduril's own canonical name (or a non-strict alias) never actually appears."""
        edges = [
            {
                "from": "Anduril",
                "to": "Elbit Systems",
                "kind": "m_and_a",
                "n": 14,
                "program": "Sigma 155 howitzer system",
            }
        ]
        source_text = {
            14: "Elbit appoints a new regional head. The old blacksmith's anvil rang loudly at the ceremony."
        }
        kept, dropped = _verify_relationship_edges(edges, source_text)
        assert kept == []
        assert len(dropped) == 1
        assert "Elbit Systems" in dropped[0]
        assert "[14]" in dropped[0]

    def test_name_in_text_strict_alias_alone_is_not_enough(self):
        text = "Elbit appoints a new regional head. The old blacksmith's anvil rang loudly."
        assert _name_in_text("Anduril", text) is False

    def test_name_in_text_strict_alias_counts_when_canonical_name_also_present(self):
        text = "Anduril unveiled its new Anvil interceptor system today."
        assert _name_in_text("Anduril", text) is True

    def test_duplicate_canonical_pair_across_spellings_is_deduped(self):
        """ "Elbit" and "Elbit Systems" are two raw spellings of the same real company -- two edges
        built off the same underlying event must collapse to one kept row, not render as a
        literal duplicate (round_3_judge.md D8 #4: "it's still there, live, now duplicated")."""
        edges = [
            {"from": "Elbit", "to": "IAI", "kind": "partnership", "n": 5},
            {"from": "Elbit Systems", "to": "IAI", "kind": "partnership", "n": 5},
        ]
        source_text = {5: "Elbit Systems and IAI announce a new joint partnership."}
        kept, dropped = _verify_relationship_edges(edges, source_text)
        assert len(kept) == 1
        assert len(dropped) == 1
        assert "כפילות" in dropped[0]

    def test_distinct_kinds_are_not_deduped(self):
        edges = [
            {"from": "Elbit", "to": "IAI", "kind": "partnership", "n": 5},
            {"from": "Elbit Systems", "to": "IAI", "kind": "contract_award", "n": 5},
        ]
        source_text = {5: "Elbit Systems and IAI announce a new joint partnership and contract."}
        kept, dropped = _verify_relationship_edges(edges, source_text)
        assert len(kept) == 2
        assert dropped == []


# --------------------------------------------------------------------------
# item 5: priority + confidence on business implications
# --------------------------------------------------------------------------


class TestPatentBizActionSchema:
    def test_accepts_valid_priority_and_confidence(self):
        action = PatentBizAction(action_he="פעולה.", priority="high", confidence=0.8, rationale_he="נימוק.")
        assert action.priority == "high"
        assert action.confidence == 0.8

    def test_defaults_when_omitted(self):
        action = PatentBizAction(action_he="פעולה.", rationale_he="נימוק.")
        assert action.priority == "medium"
        assert action.confidence == 0.5

    def test_rejects_invalid_priority(self):
        with pytest.raises(ValidationError):
            PatentBizAction(action_he="פעולה.", priority="urgent", rationale_he="נימוק.")

    @pytest.mark.parametrize("confidence", [-0.1, 1.1])
    def test_rejects_out_of_range_confidence(self, confidence):
        with pytest.raises(ValidationError):
            PatentBizAction(action_he="פעולה.", confidence=confidence, rationale_he="נימוק.")


class TestBusinessImplicationsTable:
    def test_none_for_empty_actions(self):
        assert business_implications_table([]) is None

    def test_sorted_by_priority_then_confidence(self):
        actions = [
            PatentBizAction(action_he="נמוכה.", priority="low", confidence=0.9, rationale_he="נ."),
            PatentBizAction(
                action_he="גבוהה נמוכת-ביטחון.", priority="high", confidence=0.4, rationale_he="נ."
            ),
            PatentBizAction(
                action_he="גבוהה גבוהת-ביטחון.", priority="high", confidence=0.9, rationale_he="נ."
            ),
        ]
        table = business_implications_table(actions)
        assert table is not None
        assert table["headers"] == ["עדיפות", "ביטחון", "פעולה", "נימוק"]
        actions_col = [row[2] for row in table["rows"]]
        assert actions_col == ["גבוהה גבוהת-ביטחון.", "גבוהה נמוכת-ביטחון.", "נמוכה."]

    def test_confidence_rendered_as_percentage(self):
        actions = [PatentBizAction(action_he="פעולה.", priority="medium", confidence=0.73, rationale_he="נ.")]
        table = business_implications_table(actions)
        assert table["rows"][0][1] == "73%"


class TestPriorityConfidencePrefixHe:
    def test_prefix_carries_priority_label_and_percentage(self):
        action = PatentBizAction(action_he="פעולה.", priority="high", confidence=0.9, rationale_he="נ.")
        prefix = _priority_confidence_prefix_he(action)
        assert "גבוהה" in prefix
        assert "90%" in prefix


# --------------------------------------------------------------------------
# item 6: CPC x assignee white-space matrix
# --------------------------------------------------------------------------


class TestCpcAssigneeMatrix:
    def test_counts_co_occurrences(self):
        rows = [
            {"cpc": ["G01J"], "assignees": ["Anduril"]},
            {"cpc": ["G01J"], "assignees": ["Anduril"]},
            {"cpc": ["H04N"], "assignees": ["Elbit"]},
        ]
        matrix = _cpc_assignee_matrix(rows, ["G01J", "H04N"], ["Anduril", "Elbit"])
        assert matrix == [[2, 0], [0, 1]]

    def test_zero_cell_is_a_true_white_space_gap(self):
        rows = [{"cpc": ["G01J"], "assignees": ["Anduril"]}]
        matrix = _cpc_assignee_matrix(rows, ["G01J"], ["Anduril", "Elbit"])
        assert matrix == [[1, 0]]

    def test_empty_axes_return_empty_matrix(self):
        assert _cpc_assignee_matrix([{"cpc": ["G01J"], "assignees": ["Anduril"]}], [], []) == []
