"""Unit tests for eoa.report.docx_builder: RTL correctness, structure, hyperlinks, run splitting."""

from __future__ import annotations

import datetime as dt

import docx
import pytest
from docx.oxml.ns import qn

from eoa.llm.schemas.analysis import (
    DailyReportDraft,
    OutlookIndicator,
    ReportSection,
    Sentence,
    StructuredSection,
)
from eoa.llm.schemas.reports import MonthlyReportDraft
from eoa.report import docx_builder as db
from eoa.report.qa_citations import QAResult

# --------------------------------------------------------------------------
# fixtures
# --------------------------------------------------------------------------


@pytest.fixture
def fixture_items() -> list[dict]:
    return [
        {
            "id": 101,
            "n": 1,
            "title": "Elbit Systems wins targeting pod contract",
            "source_name": "Defense News",
            "url": "https://example.com/news/elbit-pod-contract",
            "published_at": dt.date(2026, 9, 1),
            "level": "red",
            "domain": "airborne_pods",
            "summary_he": "אלביט מערכות (Elbit Systems) זכתה בחוזה לאספקת פודי כיוון (Targeting Pods) בהיקף 80 מיליון דולר.",
            "so_what_he": "להערכתנו החוזה מחזק את מעמדה התחרותי של החברה בשוק ה-EO/IR.",
            "key_facts": ["היקף החוזה: 80 מיליון דולר", "תאריך: ספטמבר 2026"],
        },
        {
            "id": 102,
            "n": 2,
            "title": "Rafael launches new C-UAS system",
            "source_name": "Janes",
            "url": "https://example.com/news/rafael-cuas",
            "published_at": dt.date(2026, 9, 2),
            "level": "orange",
            "domain": "c_uas",
            "summary_he": 'רפאל השיקה מערכת נגד כטב"מים (C-UAS) חדשה המבוססת על גילוי EO/IR.',
            "so_what_he": "המערכת צפויה להתחרות במוצרים קיימים בשוק.",
            "key_facts": ["השקה: ספטמבר 2026"],
        },
        {
            "id": 103,
            "n": 3,
            "title": "Naval EO director sea trial completed",
            "source_name": "Naval Today",
            "url": "https://example.com/news/naval-eo-director",
            "published_at": dt.date(2026, 9, 3),
            "level": "orange",
            "domain": "naval_surveillance",
            "summary_he": "בוצע ניסוי ימי (sea trial) מוצלח למערכת EO ימית חדשה.",
            "so_what_he": "מדובר בצעד משמעותי לקראת הבשלה מבצעית (operational maturity).",
            "key_facts": ["ניסוי הושלם בהצלחה"],
        },
    ]


@pytest.fixture
def fixture_events() -> list[dict]:
    return [
        {
            "id": 501,
            "item_id": 101,
            "kind": "contract_award",
            "title": "Targeting pod contract",
            "date": dt.date(2026, 9, 1),
            "amount_usd": 80_000_000,
            "currency": "USD",
            "parties": ["Elbit Systems", "USAF"],
            "customer": "USAF",
            "program": None,
            "source_name": "Defense News",
        },
    ]


@pytest.fixture
def fixture_draft() -> DailyReportDraft:
    """Goal 1 (2026-09-06): the structured sentence-per-claim shape. Citation counts are kept
    equivalent to the pre-goal-1 fixture (3 exec-summary markers [1][2][3] + 1 per section = 5
    internal citations total) so the downstream structural assertions below don't need to change."""
    return DailyReportDraft(
        exec_summary=[
            Sentence(
                text_he="אלביט מערכות (Elbit Systems) זכתה בחוזה של 80 מיליון דולר לאספקת פודי כיוון.",
                cites=[1],
            ),
            Sentence(text_he="רפאל השיקה מערכת C-UAS חדשה.", cites=[2]),
            Sentence(text_he="בוצע ניסוי ימי מוצלח למערכת EO ימית.", cites=[3]),
        ],
        sections=[
            StructuredSection(
                title_he='פודים ומטע"דים אוויריים',
                domain="airborne_pods",
                sentences=[
                    Sentence(
                        text_he="אלביט מערכות זכתה בחוזה בהיקף 80 מיליון דולר לאספקת Targeting Pods.",
                        cites=[1],
                    )
                ],
            ),
            StructuredSection(
                title_he='נגד כטב"מים',
                domain="c_uas",
                sentences=[
                    Sentence(text_he="רפאל השיקה מערכת C-UAS חדשה המבוססת על חיישני EO/IR.", cites=[2])
                ],
            ),
        ],
        outlook=[
            OutlookIndicator(text_he="להערכתנו מגמת ההשקות תימשך ברבעון הקרוב.", cites=[], is_assessment=True)
        ],
        open_points_he=["האם ידוע מי היו המתחרות שהפסידו במכרז?"],
    )


# --------------------------------------------------------------------------
# split_runs / split_runs_with_citations
# --------------------------------------------------------------------------


def test_split_runs_pure_hebrew():
    runs = db.split_runs("שלום עולם")
    assert runs == [("he", "שלום עולם")]


def test_split_runs_pure_english():
    runs = db.split_runs("Hello World")
    assert runs == [("other", "Hello World")]


def test_split_runs_mixed_hebrew_english():
    runs = db.split_runs("שלום Hello עולם")
    classes = [c for c, _ in runs]
    assert "he" in classes
    assert "other" in classes
    # Reassembling all chunks must reproduce the original text exactly.
    assert "".join(chunk for _, chunk in runs) == "שלום Hello עולם"


def test_split_runs_trailing_space_stays_with_hebrew_word():
    runs = db.split_runs("שלום Hello")
    assert runs[0][0] == "he"
    assert runs[0][1] == "שלום "
    assert runs[1] == ("other", "Hello")


def test_split_runs_with_citations_isolates_citation_token():
    tokens = db.split_runs_with_citations("משפט עם הפניה [3] להמשך.")
    classes = [c for c, _ in tokens]
    assert "cite" in classes
    cite_chunks = [chunk for c, chunk in tokens if c == "cite"]
    assert cite_chunks == ["[3]"]


def test_split_runs_bracket_pair_stays_symmetric():
    """Q5-4 (docs/qa/findings_Q5_r1.md): a parenthetical after Hebrew text used to put the opening
    '(' in the Hebrew run (it follows Hebrew) but the closing ')' in the Latin run (it follows the
    English words) -- an asymmetric split. Both brackets must land in the same (Hebrew/RTL) run.
    Uses the exact heading from the finding (gershayim ״, U+05F4, is itself a Hebrew-range char)."""
    text = "פודים ומטע״דים אוויריים (Airborne Pods & Payloads)"
    runs = db.split_runs(text)
    assert runs == [
        ("he", "פודים ומטע״דים אוויריים ("),
        ("other", "Airborne Pods & Payloads"),
        ("he", ")"),
    ]
    assert "".join(chunk for _, chunk in runs) == text


def test_split_runs_bracket_pair_symmetric_when_opened_in_latin_run():
    """The mirror case: a parenthetical opened after Latin text and closed after Hebrew text keeps
    both brackets in the Latin/other run."""
    runs = db.split_runs("Elbit (מערכת אלרון) ready")
    assert runs == [
        ("other", "Elbit ("),
        ("he", "מערכת אלרון"),
        ("other", ") ready"),
    ]


def test_bidi_html_renders_symmetric_bracket_pair():
    """The HTML export (`_bidi_html`, used for md/html reports) shares `split_runs` -- verify the
    fix actually produces the bidi-safe markup the finding asked for: both parens outside the
    <bdi dir="ltr"> span, not the closing one trailing inside it."""
    html_out = db._bidi_html("פודים ומטע״דים אוויריים (Airborne Pods & Payloads)")
    assert '(<bdi dir="ltr">Airborne Pods &amp; Payloads</bdi>)' in html_out


# --------------------------------------------------------------------------
# add_mixed_paragraph
# --------------------------------------------------------------------------


def test_add_mixed_paragraph_splits_runs_with_rtl_flags():
    doc = docx.Document()
    paragraph = db.add_mixed_paragraph(doc, "שלום Hello עולם")
    assert len(paragraph.runs) >= 2
    saw_hebrew_rtl = False
    saw_other_ltr = False
    for run in paragraph.runs:
        rpr = run._element.find(qn("w:rPr"))
        assert rpr is not None
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts is not None
        cs_font = rfonts.get(qn("w:cs"))
        if any(db._char_class(ch) == "he" for ch in run.text):
            assert run.font.rtl is True
            assert cs_font == db.HEBREW_FONT
            saw_hebrew_rtl = True
        else:
            assert not run.font.rtl
            saw_other_ltr = True
    assert saw_hebrew_rtl
    assert saw_other_ltr


def test_add_mixed_paragraph_is_bidi_and_right_aligned():
    doc = docx.Document()
    paragraph = db.add_mixed_paragraph(doc, "שלום עולם")
    ppr = paragraph._p.find(qn("w:pPr"))
    assert ppr is not None
    assert ppr.find(qn("w:bidi")) is not None
    jc = ppr.find(qn("w:jc"))
    assert jc is not None
    assert jc.get(qn("w:val")) == "right"


def test_add_mixed_paragraph_citation_is_superscript():
    """F23: `[7]` is now an internal-hyperlink run (jumping to the `src_7` appendix bookmark), so
    its run lives inside `paragraph.hyperlinks`, not `paragraph.runs` directly."""
    doc = docx.Document()
    paragraph = db.add_mixed_paragraph(doc, "עובדה חשובה [7].")
    cite_hyperlinks = [h for h in paragraph.hyperlinks if h.fragment == "src_7"]
    assert cite_hyperlinks
    cite_runs = cite_hyperlinks[0].runs
    assert cite_runs
    assert cite_runs[0].text == "[7]"
    assert cite_runs[0].font.superscript is True


# --------------------------------------------------------------------------
# add_hyperlink
# --------------------------------------------------------------------------


def test_add_hyperlink_creates_relationship_and_visible_text():
    doc = docx.Document()
    paragraph = doc.add_paragraph()
    db.add_hyperlink(paragraph, "https://example.com/x", "https://example.com/x")
    assert len(paragraph.hyperlinks) == 1
    assert paragraph.hyperlinks[0].address == "https://example.com/x"
    assert paragraph.text == "https://example.com/x"


# --------------------------------------------------------------------------
# build_docx: full structural assertions
# --------------------------------------------------------------------------


def _all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _count_hyperlinks(document) -> int:
    return sum(len(p.hyperlinks) for p in _all_paragraphs(document))


@pytest.fixture
def built_doc(fixture_draft, fixture_items, fixture_events):
    return db.build_docx(
        fixture_draft,
        fixture_items,
        fixture_events,
        period_end=dt.date(2026, 9, 4),
        qa=QAResult(passed=True),
    )


def test_build_docx_document_defaults_have_bidi_and_hebrew_font(built_doc):
    styles_elm = built_doc.styles.element
    doc_defaults = styles_elm.find(qn("w:docDefaults"))
    assert doc_defaults is not None
    ppr_default = doc_defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr"))
    assert ppr_default is not None
    assert ppr_default.find(qn("w:bidi")) is not None
    rpr_default = doc_defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
    assert rpr_default is not None
    assert rpr_default.find(qn("w:rtl")) is not None
    rfonts = rpr_default.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:cs")) == db.HEBREW_FONT

    normal = built_doc.styles["Normal"]
    normal_ppr = normal.element.find(qn("w:pPr"))
    assert normal_ppr is not None
    assert normal_ppr.find(qn("w:bidi")) is not None


def test_build_docx_headings_exist(built_doc):
    heading_texts = {
        p.text for p in built_doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    }
    assert "תקציר מנהלים" in heading_texts
    assert "טבלת אירועים עסקיים" in heading_texts
    assert "נקודות פתוחות" in heading_texts
    assert "מבט קדימה" in heading_texts
    assert "נספח מקורות" in heading_texts
    assert any("פודים" in t for t in heading_texts)
    assert any("כטב" in t for t in heading_texts)


def test_build_docx_title_page_present(built_doc):
    all_text = "\n".join(p.text for p in built_doc.paragraphs)
    assert db.TITLE_TEXT in all_text
    assert "נוצר אוטומטית" in all_text


def test_build_docx_hyperlink_count_matches_items_and_citations(built_doc, fixture_items):
    """F23: every appendix row still gets its external hyperlink, and every ``[n]`` citation in
    the prose (exec summary [1][2][3] + the two section paragraphs [1][2] = 5 in the fixture) is
    now ALSO a (internal, anchor-based) hyperlink rather than inert superscript text."""
    external = len(fixture_items)
    internal_citations = 5  # [1][2][3] in exec summary + [1] and [2] in the two sections
    assert _count_hyperlinks(built_doc) == external + internal_citations


def test_build_docx_citation_is_internal_hyperlink_to_appendix_bookmark(built_doc):
    """F23: a `[n]` citation marker is a real internal hyperlink (``w:anchor``) to the bookmark on
    its row in the sources appendix, not just superscript text."""
    citation_hyperlinks = [
        h for p in _all_paragraphs(built_doc) for h in p.hyperlinks if h.fragment.startswith("src_")
    ]
    assert citation_hyperlinks
    fragments = {h.fragment for h in citation_hyperlinks}
    assert "src_1" in fragments
    assert "src_2" in fragments
    for h in citation_hyperlinks:
        assert h.runs
        assert h.runs[0].font.superscript is True


def test_build_docx_sources_appendix_rows_are_bookmarked(built_doc):
    """F23: every numbered row in the sources appendix carries a `src_{n}` bookmark so citation
    hyperlinks (:func:`eoa.report.docx_builder.add_citation_run`) have somewhere to land."""
    xml = built_doc.element.xml
    assert 'w:name="src_1"' in xml
    assert 'w:name="src_2"' in xml
    assert 'w:name="src_3"' in xml


def test_build_docx_has_tables(built_doc):
    # events table + sources appendix table
    assert len(built_doc.tables) == 2
    for table in built_doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:bidiVisual")) is not None


def test_build_docx_events_table_has_expected_columns(built_doc):
    events_table = built_doc.tables[0]
    header_cells = [c.text for c in events_table.rows[0].cells]
    assert header_cells == ["תאריך", "סוג", "צדדים", "לקוח/תוכנית", "סכום", "מקור"]
    assert len(events_table.rows) == 2  # header + 1 event


def test_build_docx_sources_appendix_has_all_items(built_doc, fixture_items):
    appendix = built_doc.tables[-1]
    assert len(appendix.rows) == len(fixture_items) + 1  # header + N items
    header_cells = [c.text for c in appendix.rows[0].cells]
    assert header_cells == ["#", "כותרת", "מקור", "תאריך", "קישור"]


def test_build_docx_no_toc_for_daily(built_doc):
    """F10: the daily report (``include_toc`` defaults to ``False``) drops the TOC entirely rather
    than show a ``TOC`` field placeholder that displays nothing until the reader manually updates
    fields in Word."""
    xml = built_doc.element.xml
    assert "fldSimple" not in xml
    assert "יש לעדכן שדות" not in xml
    heading_texts = {
        p.text for p in built_doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    }
    assert "תוכן עניינים" not in heading_texts


def test_build_docx_real_toc_for_weekly_monthly(fixture_draft, fixture_items, fixture_events):
    """F10: ``include_toc=True`` (used by weekly/monthly) renders a real, immediately-clickable
    bookmark-based table of contents instead of a stale ``TOC`` field."""
    doc = db.build_docx(
        fixture_draft,
        fixture_items,
        fixture_events,
        period_end=dt.date(2026, 9, 4),
        qa=QAResult(passed=True),
        include_toc=True,
    )
    xml = doc.element.xml
    assert "fldSimple" not in xml
    assert "bookmarkStart" in xml
    assert "w:anchor" in xml
    heading_texts = {p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"}
    assert "תוכן עניינים" in heading_texts
    # the TOC lists every real section, including the domain sections from the fixture draft
    toc_paragraphs = [p for p in doc.paragraphs if p.style is not None and p.style.name == "List Bullet"]
    toc_texts = {p.text for p in toc_paragraphs}
    assert any("פודים" in t for t in toc_texts)
    assert any("כטב" in t for t in toc_texts)


def test_build_docx_page_field_in_footer(built_doc):
    footer_xml = built_doc.sections[0].footer._element.xml
    assert "PAGE" in footer_xml


def test_build_docx_update_fields_flagged(built_doc):
    settings_elm = built_doc.settings.element
    el = settings_elm.find(qn("w:updateFields"))
    assert el is not None
    assert el.get(qn("w:val")) == "true"


def test_build_docx_no_qa_warning_banner_for_structured_daily_draft(
    fixture_draft, fixture_items, fixture_events
):
    """Goal 1 (2026-09-06): the blanket bold-red QA-failure banner is legacy-draft-only now -- a
    structured daily draft never reaches the renderer with `qa.passed=False` and full content
    (two failures replace the whole narrative with `system_note_he` instead, at the
    ``eoa.report.daily.build_daily`` orchestration level, not here)."""
    qa = QAResult(passed=False, errors=["דוגמה לשגיאה"], bad_refs=[99])
    doc = db.build_docx(fixture_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4), qa=qa)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "אזהרה" not in all_text


def test_build_docx_qa_warning_banner_for_legacy_draft_when_failed(fixture_items, fixture_events):
    """The legacy free-prose shape (weekly/monthly/bd_territory) still degrades by silently
    stripping flagged sentences, so it still needs the visible banner."""
    legacy_draft = MonthlyReportDraft(
        exec_summary_he="תקציר [1].",
        sections=[ReportSection(title_he="סעיף", domain="c_uas", prose_he="תוכן [1].")],
        outlook_he="",
        open_points_he=[],
    )
    qa = QAResult(passed=False, errors=["דוגמה לשגיאה"], uncited_sentences=["משפט"], bad_refs=[99])
    doc = db.build_docx(legacy_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4), qa=qa)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "אזהרה" in all_text


# --------------------------------------------------------------------------
# save / validate round-trip
# --------------------------------------------------------------------------


def test_save_and_validate_docx_round_trip(built_doc, tmp_path):
    out_path = tmp_path / "sample_daily.docx"
    saved = db.save_docx(built_doc, out_path)
    assert saved == out_path
    assert out_path.exists()
    db.validate_docx(out_path)  # must not raise

    reopened = docx.Document(str(out_path))
    assert len(reopened.tables) == 2


def test_validate_docx_detects_duplicate_zip_parts(built_doc, tmp_path):
    import zipfile

    out_path = tmp_path / "broken.docx"
    db.save_docx(built_doc, out_path)

    broken_path = tmp_path / "broken_dup.docx"
    with zipfile.ZipFile(out_path) as src, zipfile.ZipFile(broken_path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            dst.writestr(item, data)
            if item.filename.endswith("document.xml"):
                dst.writestr(item, data)  # duplicate the same part on purpose

    with pytest.raises(ValueError, match="duplicate"):
        db.validate_docx(broken_path)


def test_validate_docx_detects_malformed_xml(built_doc, tmp_path):
    import zipfile

    out_path = tmp_path / "broken2.docx"
    db.save_docx(built_doc, out_path)

    broken_path = tmp_path / "broken2_malformed.docx"
    with zipfile.ZipFile(out_path) as src, zipfile.ZipFile(broken_path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith("document.xml"):
                data = data[:-20]  # truncate to corrupt the XML
            dst.writestr(item, data)

    with pytest.raises(ValueError, match="malformed"):
        db.validate_docx(broken_path)


# --------------------------------------------------------------------------
# render_markdown / render_html
# --------------------------------------------------------------------------


def test_render_markdown_contains_citations_and_appendix(fixture_draft, fixture_items, fixture_events):
    md = db.render_markdown(fixture_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4))
    assert "[1]" in md
    assert "נספח מקורות" in md
    assert "https://example.com/news/elbit-pod-contract" in md


def test_render_markdown_citations_are_links_to_appendix_anchors(
    fixture_draft, fixture_items, fixture_events
):
    """F23: `[n]` in markdown body text becomes a real `[n](#src-n)` link, and the appendix row
    carries a matching `<a id="src-n">` anchor for it to land on."""
    md = db.render_markdown(fixture_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4))
    assert "[1](#src-1)" in md
    assert "[2](#src-2)" in md
    assert '<a id="src-1"></a>' in md


def test_render_html_is_rtl_and_links_citations(fixture_draft, fixture_items, fixture_events):
    html_out = db.render_html(fixture_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4))
    assert 'dir="rtl"' in html_out
    assert 'id="src-1"' in html_out
    assert 'href="#src-1"' in html_out


def test_render_html_stylesheet_is_theme_aware(fixture_draft, fixture_items, fixture_events):
    """F21: the embedded HTML stylesheet must be theme-aware (CSS variables + a dark-mode media
    query + an explicit data-theme override) instead of a fixed white background that clashes with
    the dark Morning screen (web/src/components/reports/ReportBody.tsx embeds this HTML inline via
    `dangerouslySetInnerHTML`, not an <iframe> -- see docx_builder._EOA_HTML_STYLE's docstring)."""
    html_out = db.render_html(fixture_draft, fixture_items, fixture_events, period_end=dt.date(2026, 9, 4))

    assert "--eoa-" in html_out  # CSS custom properties present
    assert "@media (prefers-color-scheme: dark)" in html_out
    assert 'data-theme="dark"' in html_out
    assert 'data-theme="light"' in html_out
    # The root .eoa-report rule must inherit the embedding page's colours, not hard-code a light
    # background that clashes with a dark host page.
    assert "color:inherit;background:transparent" in html_out.replace(" ", "")
    # No unscoped hard-coded white body background left anywhere in the stylesheet -- the only
    # remaining "#fff" is inside the `@media print` block, guarded to `.eoa-report`, never `body`.
    assert "body{background:#fff" not in html_out.replace(" ", "")
    assert "body{background: #fff" not in html_out


def test_render_html_escapes_content():
    draft = DailyReportDraft(
        exec_summary=[Sentence(text_he="<script>alert(1)</script>", cites=[1])],
        sections=[],
        outlook=[],
        open_points_he=[],
    )
    items = [{"n": 1, "title": "t", "source_name": "s", "url": "https://x", "published_at": None}]
    html_out = db.render_html(draft, items, [])
    assert "<script>alert(1)</script>" not in html_out
    assert "&lt;script&gt;" in html_out


# --------------------------------------------------------------------------
# hebrew_date_str / fmt_date / fmt_amount
# --------------------------------------------------------------------------


def test_hebrew_date_str_format():
    s = db.hebrew_date_str(dt.date(2026, 9, 4))
    assert "2026" in s
    assert "בספטמבר" in s


def test_fmt_date_variants():
    assert db.fmt_date(None) == "—"
    assert db.fmt_date(dt.date(2026, 1, 5)) == "2026-01-05"
    assert db.fmt_date(dt.datetime(2026, 1, 5, 10, 30)) == "2026-01-05"


def test_fmt_amount_variants():
    assert db.fmt_amount({"amount_usd": None}) == "—"
    assert db.fmt_amount({"amount_usd": 80_000_000, "currency": "USD"}) == "80,000,000 USD"
