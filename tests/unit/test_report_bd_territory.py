"""Unit tests for eoa.report.bd_territory (A11 "דוח מיקוד לפיתוח עסקי, מכירה ושיווק לפי
טריטוריה"), with every DB- and LLM-touching collector monkeypatched so these run with no
Postgres and no Ollama.

Round 3 (2026-09-06): migrated alongside eoa.report.bd_territory's own structured-schema
migration (eoa.llm.schemas.bd_territory.BdTerritoryReportDraft/BdRecommendedAction, Sentence-based
citations-by-construction) -- see tests/unit/test_bd_structured_round3.py for the new round-3
scenarios (happy path renders every [n], an unknown-[n] draft is rejected then repaired, two
failures produce the deterministic substitute summary, a no-activity territory still emits the
marker, textnorm is applied, the our_company perspective rule survives in the prompt).

Run with: ``PYTHONPATH=agent python -m pytest tests/unit/test_report_bd_territory.py -q``
"""

from __future__ import annotations

import datetime as dt

import docx
import pytest

from eoa.llm.schemas.analysis import Sentence
from eoa.llm.schemas.bd_territory import BdRecommendedAction, BdTerritoryReportDraft
from eoa.report import bd_territory as bdt

MARKET_ITEMS = [
    {
        "id": 501,
        "n": 1,
        "title": "US Army awards new targeting pod contract",
        "domain": "airborne_pods",
        "source_name": "Defense News",
        "url": "https://example.com/501",
        "published_at": dt.date(2026, 8, 20),
        "level": "red",
        "geography": "US",
        "entities_mentioned": ["Elbit"],
        "summary_he": 'צבא ארה"ב הכריז על מכרז חדש לפוד כיוון.',
        "so_what_he": "הזדמנות לספקי EO/IR.",
    },
    {
        "id": 502,
        "n": 2,
        "title": "New C-UAS program announced",
        "domain": "c_uas",
        "source_name": "Janes",
        "url": "https://example.com/502",
        "published_at": dt.date(2026, 8, 25),
        "level": "orange",
        "geography": "US",
        "entities_mentioned": [],
        "summary_he": 'הוכרזה תוכנית חדשה נגד כטב"מים בארה"ב.',
        "so_what_he": "שוק צומח.",
    },
]

PLATFORM_EVENTS = [
    {
        "item_id": 999,
        "item_url": "https://example.com/999",
        "item_title": "Vendor X wins gimbal deal",
        "source_name": "Defense News",
        "published_at": dt.date(2026, 8, 22),
        "date": dt.date(2026, 8, 22),
        "platform_he": 'כטב"ם MALE',
        "payload_need_he": "מטע\"ד ג'ימבלי EO/IR",
        "buyer": "US Army",
        "vendor": "Vendor X",
        "amount_usd": 50_000_000,
        "currency": "USD",
    }
]

TENDERS_DATA = {
    "tenders": [
        {
            "title": "RFI: naval EO director",
            "agency": "US Navy",
            "country": "US",
            "deadline": dt.date(2026, 12, 1),
            "status": "open",
            "url": "https://example.gov/rfi/1",
            "published_at": dt.date(2026, 8, 1),
        }
    ],
    "forecasts": [
        {
            "platform": 'כטב"ם MALE',
            "buyer_country": "US",
            "payload_need": "מטע\"ד ג'ימבלי EO/IR",
            "likelihood": 0.6,
            "window_from": dt.date(2027, 1, 1),
            "window_to": dt.date(2027, 12, 1),
            "created_at": dt.date(2026, 8, 15),
        }
    ],
}

COMPETITORS = [
    {
        "entity_id": 1,
        "name": "Elbit",
        "country": "IL",
        "mentions": 3,
        "is_watchlist": True,
        "is_israeli_industry": True,
        "recent_wins": [
            {
                "id": 1,
                "item_id": 501,
                "title": "Targeting pod win",
                "date": dt.date(2026, 8, 20),
                "amount_usd": 10_000_000,
                "currency": "USD",
                "customer": "US Army",
                "program": None,
            }
        ],
    }
]

CONFERENCES_DATA = {
    "territory": [
        {
            "name": "AUSA",
            "start_date": dt.date(2026, 10, 10),
            "end_date": dt.date(2026, 10, 12),
            "city": "Washington",
            "relevance": 5,
            "status": "confirmed",
            "organizer": "Association of the United States Army",
        }
    ],
    "international": [],
}


def _draft_fixture() -> BdTerritoryReportDraft:
    return BdTerritoryReportDraft(
        exec_summary=[
            Sentence(text_he='צבא ארה"ב הכריז על מכרז חדש לפוד כיוון.', cites=[1]),
            Sentence(text_he='שוק ה-C-UAS בארה"ב צומח.', cites=[2]),
        ],
        market_bullets=[
            Sentence(text_he='צבא ארה"ב מקדם מכרז לפוד כיוון חדש.', cites=[1]),
            Sentence(text_he='תוכנית C-UAS חדשה הוכרזה בארה"ב.', cites=[2]),
            Sentence(text_he="Vendor X זכתה בעסקת ג'ימבל בהיקף 50 מיליון דולר.", cites=[3]),
        ],
        competitor_moves=[
            Sentence(text_he='Elbit זכתה בעסקת פוד כיוון בארה"ב.', cites=[1]),
        ],
        sections=[],
        recommended_actions=[
            BdRecommendedAction(
                action_he="ליזום פגישת היכרות עם US Army לקראת ה-RFI הימי",
                priority="H",
                rationale=[
                    Sentence(text_he='נפתח RFI לכיוון ימי בארה"ב.', cites=[4]),
                    Sentence(text_he="תחזית הרכש תומכת בכך.", cites=[5]),
                ],
                owner_role_he="פיתוח עסקי",
                timing_he="מיידי",
                target="US Navy",
                confidence=0.7,
            ),
            BdRecommendedAction(
                action_he="להציג יכולות ג'ימבל בכנס AUSA הקרוב",
                priority="M",
                rationale=[Sentence(text_he="Vendor X כבר פעילה בשוק הזה.", cites=[3])],
                owner_role_he="שיווק",
                timing_he="רבעון הקרוב",
                target="AUSA",
                confidence=0.6,
            ),
        ],
        analyst_note_he=None,
        risks_assumptions_he="הדוח מבוסס על כיסוי מקורות חלקי בחלון הזמן שנבדק בלבד.",
        open_points_he=["האם ידוע על תקציב מאושר ל-RFI הימי?"],
    )


@pytest.fixture
def patch_bd_collectors(monkeypatch, tmp_path):
    monkeypatch.setattr(
        bdt, "collect_market_items", lambda t, s, e, max_items=250: [dict(it) for it in MARKET_ITEMS]
    )
    monkeypatch.setattr(
        bdt, "collect_platform_events", lambda t, s, e, limit=25: [dict(ev) for ev in PLATFORM_EVENTS]
    )
    monkeypatch.setattr(
        bdt,
        "collect_tenders_and_forecasts",
        lambda t, limit=20: {
            "tenders": [dict(x) for x in TENDERS_DATA["tenders"]],
            "forecasts": [dict(x) for x in TENDERS_DATA["forecasts"]],
        },
    )
    monkeypatch.setattr(
        bdt, "collect_active_competitors", lambda t, ids, s, e, limit=15: [dict(c) for c in COMPETITORS]
    )
    monkeypatch.setattr(
        bdt,
        "collect_conferences_for_territory",
        lambda t, months=12, international_limit=5: {
            "territory": [dict(c) for c in CONFERENCES_DATA["territory"]],
            "international": [],
        },
    )
    monkeypatch.setattr(bdt, "collect_dormant_watchlist_competitors", lambda t, active, limit=20: [])
    monkeypatch.setattr(bdt, "draft_bd_territory", lambda *a, **k: _draft_fixture())
    monkeypatch.setattr(bdt, "_persist_report", lambda *a, **k: 777)
    monkeypatch.setattr(bdt, "_report_path", lambda territory, period_end, ext: tmp_path / f"bd.{ext}")
    return tmp_path


def test_lookback_range_defaults_to_90_days_window():
    start, end = bdt.lookback_range(90, dt.date(2026, 9, 6))
    assert end == dt.date(2026, 9, 6)
    assert (end - start).days == 89


def test_territory_label_normalizes():
    assert bdt.territory_label("us") == "US"
    assert bdt.territory_label("uk") == "GB"


def test_render_sentences_appends_citation_markers():
    text = bdt._render_sentences([Sentence(text_he="בולט בלי נקודה", cites=[1])])
    assert text == "בולט בלי נקודה [1]"


def test_sentences_bullets_text_one_per_line():
    text = bdt._sentences_bullets_text(
        [Sentence(text_he="בולט ראשון", cites=[1]), Sentence(text_he="בולט שני", cites=[2])]
    )
    lines = text.split("\n")
    assert lines == ["בולט ראשון [1]", "בולט שני [2]"]


def test_format_market_items_block_empty():
    assert "לא זוהו" in bdt.format_market_items_block([])


def test_format_competitors_block_empty():
    assert "לא זוהו מתחרים" in bdt.format_competitors_block([])


def test_build_bd_territory_qa_passes_on_fixture_draft(patch_bd_collectors):
    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    assert paths.qa.passed, paths.qa.errors
    assert paths.report_id == 777
    assert paths.territory == "US"


def test_build_bd_territory_renders_expected_tables(patch_bd_collectors):
    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    doc = docx.Document(str(paths.docx))
    table_headers = [[c.text for c in t.rows[0].cells] for t in doc.tables]

    assert ["תאריך", "פלטפורמה/תוכנית", "רוכש", "ספק", "סכום", "צורך EO/IR נגזר", "מקור"] in table_headers
    assert ["כותרת", "גורם מזמין", "דדליין", "סטטוס", "קישור"] in table_headers
    # Round 5 P6 (B2): competitors_table gained a trailing "דרג" (tier A/B/C) column.
    assert ["מתחרה", "מדינה", "אזכורים בחלון", "תעשייה ישראלית", "זכייה אחרונה", "דרג"] in table_headers
    assert ["שם", "תאריכים", "עיר", "סטטוס", "מארגן", "רלוונטיות"] in table_headers
    assert ["עדיפות", "פעולה", "נימוק", "אחראי", "תזמון"] in table_headers

    heading_texts = {p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"}
    assert "תמונת שוק בטריטוריה" in heading_texts
    assert "מהלכי מתחרים בטריטוריה" in heading_texts
    assert "סיכונים והנחות" in heading_texts
    assert "דוח מיקוד לפיתוח עסקי — US" in {
        p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Title"
    }


def test_build_bd_territory_actions_table_sorted_by_priority(patch_bd_collectors):
    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    doc = docx.Document(str(paths.docx))
    actions_table = next(
        t
        for t in doc.tables
        if [c.text for c in t.rows[0].cells] == ["עדיפות", "פעולה", "נימוק", "אחראי", "תזמון"]
    )
    priorities = [row.cells[0].text for row in actions_table.rows[1:]]
    assert priorities == ["גבוהה", "בינונית"]


def test_build_bd_territory_no_items_skips_llm_and_still_persists(monkeypatch, tmp_path):
    monkeypatch.setattr(bdt, "collect_market_items", lambda t, s, e, max_items=250: [])
    monkeypatch.setattr(bdt, "collect_platform_events", lambda t, s, e, limit=25: [])
    monkeypatch.setattr(
        bdt, "collect_tenders_and_forecasts", lambda t, limit=20: {"tenders": [], "forecasts": []}
    )
    monkeypatch.setattr(bdt, "collect_active_competitors", lambda t, ids, s, e, limit=15: [])
    monkeypatch.setattr(
        bdt,
        "collect_conferences_for_territory",
        lambda t, months=12, international_limit=5: {"territory": [], "international": []},
    )
    monkeypatch.setattr(bdt, "collect_dormant_watchlist_competitors", lambda t, active, limit=20: [])
    monkeypatch.setattr(bdt, "_persist_report", lambda *a, **k: 1)
    monkeypatch.setattr(bdt, "_report_path", lambda territory, period_end, ext: tmp_path / f"bd.{ext}")

    def _fail_if_called(*a, **k):
        raise AssertionError("chat_structured should not be called with zero items")

    monkeypatch.setattr("eoa.report.bd_territory.chat_structured", _fail_if_called)

    paths = bdt.build_bd_territory("IL", 90, period_end=dt.date(2026, 9, 6))
    assert paths.qa.passed


def test_extend_registry_with_source_items_appends_new_entry():
    citation_items = [{"id": 1, "n": 1}]
    entries = [{"item_id": 999, "item_title": "t", "item_url": "u", "source_name": "s", "published_at": None}]
    out = bdt._extend_registry_with_source_items(citation_items, entries)
    assert len(out) == 2
    assert entries[0]["n"] == 2


def test_extend_registry_with_source_items_reuses_existing_entry():
    citation_items = [{"id": 999, "n": 1}]
    entries = [{"item_id": 999}]
    out = bdt._extend_registry_with_source_items(citation_items, entries)
    assert len(out) == 1
    assert entries[0]["n"] == 1


def test_extend_registry_with_tenders_assigns_sequential_n():
    citation_items = [{"id": 1, "n": 1}]
    data = {"tenders": [{"title": "t1"}], "forecasts": [{"platform": "p1"}]}
    out = bdt._extend_registry_with_tenders(citation_items, data)
    assert len(out) == 3
    assert data["tenders"][0]["n"] == 2
    assert data["forecasts"][0]["n"] == 3


def test_extend_registry_with_conferences_covers_both_lists():
    citation_items = [{"id": 1, "n": 1}]
    data = {"territory": [{"name": "AUSA"}], "international": [{"name": "DSEI"}]}
    out = bdt._extend_registry_with_conferences(citation_items, data)
    assert len(out) == 3
    assert data["territory"][0]["n"] == 2
    assert data["international"][0]["n"] == 3


def test_attach_win_citations_reuses_market_item_n():
    citation_items = [{"id": 501, "n": 1}, {"id": 502, "n": 2}]
    competitors = [{"recent_wins": [{"item_id": 501}, {"item_id": 999}]}]
    bdt._attach_win_citations(citation_items, competitors)
    assert competitors[0]["recent_wins"][0]["n"] == 1
    assert competitors[0]["recent_wins"][1]["n"] is None


def test_format_tenders_block_includes_citation_numbers():
    data = {"tenders": [{"title": "RFI", "n": 4, "agency": "US Navy"}], "forecasts": []}
    block = bdt.format_tenders_block(data)
    assert "[4]" in block


def test_format_conferences_block_includes_citation_numbers():
    data = {"territory": [{"name": "AUSA", "n": 7}], "international": []}
    block = bdt.format_conferences_block(data)
    assert "[7]" in block


def test_format_competitors_block_includes_win_citation_numbers():
    competitors = [
        {
            "name": "Elbit",
            "is_israeli_industry": True,
            "country": "IL",
            "mentions": 1,
            "recent_wins": [{"title": "win", "n": 3, "date": None, "amount_usd": None}],
        }
    ]
    block = bdt.format_competitors_block(competitors)
    assert "[3]" in block


# --------------------------------------------------------------------------
# BD-1 (docs/qa/findings_Q3_r2.md) regression tests
# --------------------------------------------------------------------------


def test_conferences_table_includes_status_and_organizer():
    data = {
        "territory": [
            {
                "name": "AUSA",
                "start_date": dt.date(2026, 10, 10),
                "end_date": dt.date(2026, 10, 12),
                "city": "Washington",
                "relevance": 5,
                "status": "confirmed",
                "organizer": "Association of the United States Army",
            }
        ],
        "international": [],
    }
    table = bdt.conferences_table(data)
    assert table["headers"] == ["שם", "תאריכים", "עיר", "סטטוס", "מארגן", "רלוונטיות"]
    assert table["rows"][0] == [
        "AUSA",
        "2026-10-10 - 2026-10-12",
        "Washington",
        "מאושר",
        "Association of the United States Army",
        5,
    ]


def test_format_conferences_block_includes_status_and_organizer():
    data = {
        "territory": [
            {
                "name": "AUSA",
                "n": 1,
                "start_date": dt.date(2026, 10, 10),
                "end_date": dt.date(2026, 10, 12),
                "city": "Washington",
                "status": "estimated",
                "organizer": "AUSA Inc.",
            }
        ],
        "international": [],
    }
    block = bdt.format_conferences_block(data)
    assert "משוער" in block
    assert "AUSA Inc." in block
    assert "2026-10-10" in block and "2026-10-12" in block


def test_conference_status_he_falls_back_to_raw_value():
    assert bdt._conference_status_he({"status": "confirmed"}) == "מאושר"
    assert bdt._conference_status_he({"status": "estimated"}) == "משוער"
    assert bdt._conference_status_he({"status": None}) == "—"


def test_watchlist_competitor_names_filters_non_watchlist():
    competitors = [
        {"name": "Elbit", "is_watchlist": True},
        {"name": "Vendor X", "is_watchlist": False},
    ]
    assert bdt._watchlist_competitor_names(competitors) == {"Elbit"}


def test_action_promoted_competitor_detects_promotion_verb_and_name():
    action = BdRecommendedAction(
        action_he="להציג יכולת של Shield AI בכנס AUSA הקרוב",
        priority="M",
        rationale=[Sentence(text_he="Shield AI פעילה בשוק.", cites=[1])],
        owner_role_he="שיווק",
        timing_he="רבעון הקרוב",
    )
    assert bdt._action_promoted_competitor(action, {"Shield AI"}) == "Shield AI"


def test_action_promoted_competitor_ignores_non_watchlist_mentions():
    action = BdRecommendedAction(
        action_he="לפנות ללקוח בנוגע ל-Shield AI כמתחרה בשוק",
        priority="M",
        rationale=[Sentence(text_he="Shield AI מתחרה בשוק.", cites=[1])],
        owner_role_he="פיתוח עסקי",
        timing_he="מיידי",
    )
    # No promotion verb present -- monitoring/approaching the customer about a competitor is fine.
    assert bdt._action_promoted_competitor(action, {"Shield AI"}) is None


def test_perspective_violations_flags_competitor_promoting_action():
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])],
        recommended_actions=[
            BdRecommendedAction(
                action_he="להציג יכולת של Shield AI בכנס AUSA",
                priority="H",
                rationale=[Sentence(text_he="Shield AI פעילה בתחום.", cites=[1])],
                owner_role_he="שיווק",
                timing_he="רבעון הקרוב",
            ),
            BdRecommendedAction(
                action_he="ליזום פגישה עם הלקוח בנוגע למכרז",
                priority="M",
                rationale=[Sentence(text_he="נפתח מכרז רלוונטי.", cites=[1])],
                owner_role_he="פיתוח עסקי",
                timing_he="מיידי",
            ),
        ],
    )
    competitors = [{"name": "Shield AI", "is_watchlist": True}]
    violations = bdt._perspective_violations(draft, competitors)
    assert len(violations) == 1
    assert violations[0][1] == "Shield AI"


def test_perspective_violations_empty_when_no_watchlist_competitors():
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])],
        recommended_actions=[
            BdRecommendedAction(
                action_he="להציג יכולת של Shield AI בכנס AUSA",
                priority="H",
                rationale=[Sentence(text_he="Shield AI פעילה בתחום.", cites=[1])],
                owner_role_he="שיווק",
                timing_he="רבעון הקרוב",
            )
        ],
    )
    assert bdt._perspective_violations(draft, []) == []


def test_drop_perspective_violations_removes_only_offending_action():
    keep = BdRecommendedAction(
        action_he="ליזום פגישה עם הלקוח",
        priority="M",
        rationale=[Sentence(text_he="נפתח מכרז.", cites=[1])],
        owner_role_he="פיתוח עסקי",
        timing_he="מיידי",
    )
    drop = BdRecommendedAction(
        action_he="להציג יכולת של Shield AI בכנס AUSA",
        priority="H",
        rationale=[Sentence(text_he="Shield AI פעילה.", cites=[1])],
        owner_role_he="שיווק",
        timing_he="רבעון הקרוב",
    )
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])], recommended_actions=[keep, drop]
    )
    cleaned = bdt._drop_perspective_violations(draft, [(drop, "Shield AI")])
    assert cleaned.recommended_actions == [keep]


def test_strip_placeholder_echoes_removes_summary_sentence_with_fictional_tender():
    draft = BdTerritoryReportDraft(
        exec_summary=[
            Sentence(text_he='צבא ארה"ב מתמודד עם איומי רחפנים קטנים.', cites=[1]),
            Sentence(
                text_he="הפעולה הדחופה ביותר המומלצת היא ליזום פגישת היכרות עם גורם מזמין לקראת מכרז X.",
                cites=[1],
            ),
        ],
    )
    cleaned = bdt._strip_placeholder_echoes(draft)
    texts = [s.text_he for s in cleaned.exec_summary]
    assert not any("מכרז X" in t for t in texts)
    assert 'צבא ארה"ב מתמודד עם איומי רחפנים קטנים.' in texts


def test_strip_placeholder_echoes_removes_bullet_and_action_with_generic_names():
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])],
        market_bullets=[
            Sentence(text_he="התפתחות אמיתית.", cites=[1]),
            Sentence(text_he="התפתחות בכנס Z הקרוב.", cites=[2]),
        ],
        recommended_actions=[
            BdRecommendedAction(
                action_he="להציג יכולת Y בכנס Z הקרוב",
                priority="M",
                rationale=[Sentence(text_he="נימוק.", cites=[1])],
                owner_role_he="שיווק",
                timing_he="מיידי",
            ),
            BdRecommendedAction(
                action_he="ליזום פגישה עם US Army",
                priority="H",
                rationale=[Sentence(text_he="נימוק אמיתי.", cites=[1])],
                owner_role_he="פיתוח עסקי",
                timing_he="מיידי",
            ),
        ],
    )
    cleaned = bdt._strip_placeholder_echoes(draft)
    assert [s.text_he for s in cleaned.market_bullets] == ["התפתחות אמיתית."]
    assert len(cleaned.recommended_actions) == 1
    assert cleaned.recommended_actions[0].action_he == "ליזום פגישה עם US Army"


def test_strip_placeholder_echoes_noop_when_clean():
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר אמיתי.", cites=[1])],
        market_bullets=[Sentence(text_he="בולט אמיתי.", cites=[1])],
        recommended_actions=[
            BdRecommendedAction(
                action_he="פעולה",
                priority="M",
                rationale=[Sentence(text_he="נימוק.", cites=[1])],
                owner_role_he="מכירות",
                timing_he="מיידי",
            )
        ],
    )
    cleaned = bdt._strip_placeholder_echoes(draft)
    assert cleaned is draft


def test_cap_draft_lengths_truncates_runaway_bullets_and_actions():
    bullets = [Sentence(text_he=f"בולט מספר {i}.", cites=[1]) for i in range(15)]
    actions = [
        BdRecommendedAction(
            action_he=f"פעולה {i}",
            priority="M",
            rationale=[Sentence(text_he="נימוק.", cites=[1])],
            owner_role_he="מכירות",
            timing_he="מיידי",
        )
        for i in range(20)
    ]
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])],
        market_bullets=bullets,
        recommended_actions=actions,
    )
    capped = bdt._cap_draft_lengths(draft)
    assert len(capped.market_bullets) == 8
    assert len(capped.recommended_actions) == 8
    assert capped.market_bullets == bullets[:8]
    assert capped.recommended_actions == actions[:8]


def test_cap_draft_lengths_noop_when_within_limits():
    draft = BdTerritoryReportDraft(
        exec_summary=[Sentence(text_he="תקציר.", cites=[1])],
        market_bullets=[Sentence(text_he="בולט.", cites=[1])],
        recommended_actions=[
            BdRecommendedAction(
                action_he="פעולה",
                priority="M",
                rationale=[Sentence(text_he="נימוק.", cites=[1])],
                owner_role_he="מכירות",
                timing_he="מיידי",
            )
        ],
    )
    assert bdt._cap_draft_lengths(draft) is draft


def test_our_company_block_he_reads_config_default():
    block = bdt._our_company_block_he()
    assert "החברה שלנו" in block or "our_company" in block.lower()


def test_bd_table_counts_context_he_lists_nonzero_counts():
    counts = bdt.BdTableCounts(events=2, tenders=1, forecasts=0, competitors=3, conferences=0)
    text = counts.context_he()
    assert "2 אירועי" in text
    assert "1 מכרזים" in text
    assert "3 מתחרים" in text
    assert "תחזיות" not in text
    assert "כנסים" not in text


def test_bd_table_counts_context_he_empty_when_all_zero():
    assert bdt.BdTableCounts().context_he() == "אין (כל הטבלאות ריקות בתקופה זו)."


def test_tables_only_draft_used_when_items_empty_but_tables_present():
    draft = bdt.draft_bd_territory(
        "US",
        90,
        "",
        "",
        "",
        "",
        "",
        has_items=False,
        table_counts=bdt.BdTableCounts(competitors=3),
    )
    assert "לא זוהו פריטי שוק חדשים" in draft.system_note_he
    assert "3 מתחרים פעילים" in draft.system_note_he
    assert draft.exec_summary == []
    assert draft.recommended_actions == []


def test_no_items_draft_used_when_everything_empty():
    draft = bdt.draft_bd_territory(
        "US", 90, "", "", "", "", "", has_items=False, table_counts=bdt.BdTableCounts()
    )
    assert "אין ממצאים" in draft.system_note_he
    assert draft.exec_summary == []


def test_collect_active_competitors_excludes_zero_activity_company(monkeypatch):
    def fake_fetchall(query, params=None):
        if "FROM entities" in query and "unnest" not in query:
            return [
                {"id": 1, "name": "Active Co", "country": "US", "relevance": 0.9, "is_watchlist": True},
                {"id": 2, "name": "Dormant Co", "country": "US", "relevance": 0.9, "is_watchlist": True},
            ]
        if "unnest(entities_mentioned)" in query:
            return [{"name": "Active Co", "n": 2}]
        if "FROM events" in query:
            return []
        if "FROM graph_edges" in query:
            return [{"n": 0}]
        raise AssertionError(f"unexpected query: {query}")

    monkeypatch.setattr(bdt, "_fetchall", fake_fetchall)
    out = bdt.collect_active_competitors("US", [10, 11], dt.date(2026, 1, 1), dt.date(2026, 3, 1))
    names = [c["name"] for c in out]
    assert names == ["Active Co"]


def test_collect_dormant_watchlist_competitors_lists_names_not_in_active_set(monkeypatch):
    def fake_fetchall(query, params=None):
        return [
            {"name": "Active Co", "country": "US"},
            {"name": "Dormant Co", "country": "US"},
            {"name": "Other Territory Co", "country": "IL"},
        ]

    monkeypatch.setattr(bdt, "_fetchall", fake_fetchall)
    dormant = bdt.collect_dormant_watchlist_competitors("US", {"Active Co"})
    assert dormant == ["Dormant Co"]


def test_build_bd_territory_adds_dormant_note_when_few_competitors(patch_bd_collectors, monkeypatch):
    monkeypatch.setattr(
        bdt, "collect_active_competitors", lambda t, ids, s, e, limit=15: [dict(c) for c in COMPETITORS]
    )
    monkeypatch.setattr(
        bdt, "collect_dormant_watchlist_competitors", lambda t, active, limit=20: ["Dormant Co", "Sleepy Co"]
    )
    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    html_text = paths.html.read_text(encoding="utf-8")
    assert "Dormant Co" in html_text
    assert "Sleepy Co" in html_text


def test_build_bd_territory_no_dormant_note_when_three_or_more_competitors(patch_bd_collectors, monkeypatch):
    many_competitors = [dict(COMPETITORS[0]) for _ in range(3)]
    for idx, c in enumerate(many_competitors):
        c["name"] = f"Competitor {idx}"
        c["recent_wins"] = []
    monkeypatch.setattr(bdt, "collect_active_competitors", lambda t, ids, s, e, limit=15: many_competitors)
    monkeypatch.setattr(
        bdt, "collect_dormant_watchlist_competitors", lambda t, active, limit=20: ["Dormant Co"]
    )
    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    html_text = paths.html.read_text(encoding="utf-8")
    assert "Dormant Co" not in html_text


def test_build_bd_territory_drops_perspective_violation_after_failed_retry(patch_bd_collectors, monkeypatch):
    bad_action = BdRecommendedAction(
        action_he="להציג יכולת של Shield AI בכנס AUSA",
        priority="H",
        rationale=[Sentence(text_he="Shield AI פעילה בתחום.", cites=[1])],
        owner_role_he="שיווק",
        timing_he="רבעון הקרוב",
    )
    good_action = BdRecommendedAction(
        action_he="ליזום פגישת היכרות עם US Army",
        priority="M",
        rationale=[Sentence(text_he='נפתח RFI לכיוון ימי בארה"ב.', cites=[4])],
        owner_role_he="פיתוח עסקי",
        timing_he="מיידי",
    )
    violating_draft = _draft_fixture().model_copy(update={"recommended_actions": [bad_action, good_action]})

    monkeypatch.setattr(bdt, "draft_bd_territory", lambda *a, **k: violating_draft)
    monkeypatch.setattr(
        bdt,
        "collect_active_competitors",
        lambda t, ids, s, e, limit=15: (
            [dict(c) for c in COMPETITORS]
            + [
                {
                    "entity_id": 2,
                    "name": "Shield AI",
                    "country": "US",
                    "mentions": 1,
                    "is_watchlist": True,
                    "is_israeli_industry": False,
                    "recent_wins": [],
                }
            ]
        ),
    )
    # The regeneration retry itself still violates (simulating a stubborn model) -- the fallback
    # must drop exactly the offending action rather than persist a competitor-promoting one.
    monkeypatch.setattr(bdt, "_perspective_corrective_retry", lambda *a, **k: violating_draft)

    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    html_text = paths.html.read_text(encoding="utf-8")
    # Shield AI may legitimately appear in the competitors table (it *is* an active competitor) --
    # what must never survive is the action that recommends promoting it.
    assert "להציג יכולת של Shield AI" not in html_text
    assert "ליזום פגישת היכרות" in html_text
    assert "US Army" in html_text


# ---------------------------------------------------------------------------------------------
# D7 round-1 fix (docs/qa/loop/round_1_fixes.md, actions_table_nonempty): deterministic
# candidate-actions fallback when the LLM produced no cited actions.
# ---------------------------------------------------------------------------------------------


def test_deterministic_candidate_actions_covers_every_data_source():
    competitors = [dict(c) for c in COMPETITORS]
    competitors[0]["recent_wins"][0]["n"] = 1
    tenders_data = {
        "tenders": [{**TENDERS_DATA["tenders"][0], "n": 4}],
        "forecasts": [],
    }
    conferences_data = {"territory": [{**CONFERENCES_DATA["territory"][0], "n": 6}], "international": []}
    events = [{**PLATFORM_EVENTS[0], "n": 3}]

    actions = bdt._deterministic_candidate_actions(competitors, tenders_data, conferences_data, events)

    assert len(actions) == 4
    action_texts = [a.action_he for a in actions]
    assert any("לבחון תגובה תחרותית ל-Elbit" in t for t in action_texts)
    assert any("להיערך ל-AUSA" in t for t in action_texts)
    assert any("לבחון מענה ל-RFI: naval EO director" in t for t in action_texts)
    assert any("לפנות ל-US Army בנושא" in t for t in action_texts)
    # every action carries a citation into the registry
    for action in actions:
        assert action.rationale
        for sentence in action.rationale:
            assert sentence.cites


def test_deterministic_candidate_actions_empty_when_no_data():
    assert bdt._deterministic_candidate_actions([], {"tenders": []}, {"territory": []}, []) == []


def test_deterministic_candidate_actions_skips_rows_without_registry_n():
    # A fresh, independent fixture (not the shared COMPETITORS module list, which other tests in
    # this file mutate in place by adding an "n" key to its nested recent_wins dicts) with no "n"
    # key on the win at all.
    competitors = [
        {
            "name": "Elbit",
            "recent_wins": [
                {"item_id": 501, "title": "Targeting pod win", "program": None},
            ],
        }
    ]
    actions = bdt._deterministic_candidate_actions(competitors, {"tenders": []}, {"territory": []}, [])
    assert actions == []


def test_recommended_actions_table_deterministic_uses_alternate_title_and_note():
    draft = _draft_fixture()
    table = bdt.recommended_actions_table(draft, deterministic=True)
    assert table["title_he"] == bdt._DETERMINISTIC_ACTIONS_TITLE_HE
    assert table["note_he"] == bdt._DETERMINISTIC_ACTIONS_NOTE_HE


def test_recommended_actions_table_normal_has_no_note():
    draft = _draft_fixture()
    table = bdt.recommended_actions_table(draft)
    assert table["title_he"] == "נקודות כניסה ופעולות מומלצות"
    assert "note_he" not in table


def test_recommended_actions_table_renders_rationale_with_citation_markers():
    draft = _draft_fixture()
    table = bdt.recommended_actions_table(draft)
    rationale_cells = [row[2] for row in table["rows"]]
    assert any("[4]" in cell and "[5]" in cell for cell in rationale_cells)


def test_build_bd_territory_uses_deterministic_actions_when_llm_actions_empty(
    patch_bd_collectors, monkeypatch
):
    """When the drafted report ends up with zero recommended actions (LLM produced none, or every
    one was stripped by the citation/perspective gates), the rendered report still carries a
    populated, deterministic actions table instead of an empty section."""
    empty_actions_draft = _draft_fixture().model_copy(update={"recommended_actions": []})
    monkeypatch.setattr(bdt, "draft_bd_territory", lambda *a, **k: empty_actions_draft)

    paths = bdt.build_bd_territory("US", 90, period_end=dt.date(2026, 9, 6))
    md_text = paths.md.read_text(encoding="utf-8")  # plain markdown -- no bidi <bdi> tags splitting names

    assert bdt._DETERMINISTIC_ACTIONS_TITLE_HE in md_text
    assert bdt._DETERMINISTIC_ACTIONS_NOTE_HE in md_text
    assert "לבחון תגובה תחרותית ל-Elbit" in md_text
