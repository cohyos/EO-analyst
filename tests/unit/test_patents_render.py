"""Tests for eoa.patents.render (2026-09-06 goal: bidi isolation for Latin/English values in a
plain-Markdown table cell -- see the module docstring for why this exists alongside
eoa.report.docx_builder's own run-level bidi splitting)."""

from __future__ import annotations

import docx

from eoa.patents.render import (
    ascii_timeline,
    find_table_by_headers,
    inject_advance_footnotes_html,
    inject_advance_footnotes_md,
    insert_section_before_html_appendix,
    insert_section_before_md_appendix,
    ltr_isolate,
    ltr_isolate_if_latin,
    ltr_join,
    shade_table_cell,
    shade_timeline_table_rows,
    svg_timeline_bar_chart,
)

_LRI = "⁦"
_PDI = "⁩"


class TestLtrIsolate:
    def test_wraps_latin_value(self):
        out = ltr_isolate("US9197834B2")
        assert out == f"{_LRI}US9197834B2{_PDI}"

    def test_passthrough_for_none(self):
        assert ltr_isolate(None) == "—"

    def test_passthrough_for_placeholder_dash(self):
        assert ltr_isolate("—") == "—"

    def test_passthrough_for_empty_string(self):
        assert ltr_isolate("") == "—"


class TestLtrJoin:
    def test_joins_and_isolates_as_one_unit(self):
        out = ltr_join(["Anduril", "RTX"])
        assert out == f"{_LRI}Anduril, RTX{_PDI}"

    def test_empty_list_is_dash(self):
        assert ltr_join([]) == "—"

    def test_none_is_dash(self):
        assert ltr_join(None) == "—"

    def test_custom_separator(self):
        out = ltr_join(["G01J5", "G02B23/27"], sep=" | ")
        assert out == f"{_LRI}G01J5 | G02B23/27{_PDI}"


class TestLtrIsolateIfLatin:
    def test_wraps_pure_latin_title(self):
        out = ltr_isolate_if_latin("Digital ROIC enhancement and repetition")
        assert out.startswith(_LRI) and out.endswith(_PDI)

    def test_leaves_hebrew_text_untouched(self):
        text = "התעשייה הישראלית אינה נוכחת בנוף הפטנטים"
        assert ltr_isolate_if_latin(text) == text

    def test_leaves_mixed_hebrew_latin_untouched(self):
        """A title mixing scripts is left alone -- docx_builder's own per-run bidi splitting (which
        this helper is additive to, not a replacement for) already handles that case correctly
        inside the docx/html renderers; wrapping the *whole* mixed string LTR here would be wrong."""
        text = "אלביט מערכות (Elbit Systems) מכריזה על פטנט חדש"
        assert ltr_isolate_if_latin(text) == text

    def test_passthrough_for_none(self):
        assert ltr_isolate_if_latin(None) == "—"


class TestAsciiTimeline:
    def test_empty_dict_returns_honest_message(self):
        assert "אין נתוני" in ascii_timeline({})

    def test_renders_one_line_per_year_sorted(self):
        out = ascii_timeline({2021: 1, 2019: 5})
        lines = out.split("\n")
        assert lines[0].startswith("2019")
        assert lines[1].startswith("2021")

    def test_bar_length_scales_with_max(self):
        out = ascii_timeline({2019: 10, 2020: 5}, max_width=20)
        bar_2019 = out.split("\n")[0]
        bar_2020 = out.split("\n")[1]
        assert bar_2019.count("█") > bar_2020.count("█")


class TestSvgTimelineBarChart:
    def test_empty_dict_returns_placeholder_svg(self):
        out = svg_timeline_bar_chart({})
        assert out.startswith("<svg")
        assert "אין נתוני" in out

    def test_renders_one_rect_per_year(self):
        out = svg_timeline_bar_chart({2019: 2, 2020: 5, 2021: 1})
        assert out.count("<rect") - 1 == 3  # -1 for the background rect

    def test_is_well_formed_enough_to_close(self):
        out = svg_timeline_bar_chart({2019: 1})
        assert out.startswith("<svg") and out.rstrip().endswith("</svg>")


class TestInjectAdvanceFootnotesMd:
    def test_inserts_line_after_first_citation_only(self):
        text = (
            "## נוף הפטנטים\n\nראה [1](#src-1) לפרטים. וגם [1](#src-1) שוב.\n\n"
            "## נספח מקורות\n\n| # |\n|---|\n| [1](#src-1) |\n"
        )
        out = inject_advance_footnotes_md(text, {1: "תיאור התקדמות."})
        assert out.count("התקדמות פטנט [1]") == 1
        # inserted right after the paragraph line, before the appendix heading
        lines = out.split("\n")
        idx = next(i for i, line in enumerate(lines) if "התקדמות פטנט [1]" in line)
        assert "[1](#src-1)" in lines[idx - 1]
        assert lines.index("## נספח מקורות") > idx

    def test_no_op_for_empty_map(self):
        text = "## X\n\n[1](#src-1)\n"
        assert inject_advance_footnotes_md(text, {}) == text

    def test_ignores_citation_number_not_in_map(self):
        text = "## X\n\n[2](#src-2)\n\n## נספח מקורות\n"
        out = inject_advance_footnotes_md(text, {1: "תיאור."})
        assert "התקדמות" not in out


class TestInjectAdvanceFootnotesHtml:
    def test_inserts_paragraph_after_first_citation(self):
        text = (
            '<h2 id="sec-0">נוף הפטנטים</h2>\n<p>ראה <a href="#src-1" class="cite">[1]</a> לפרטים.</p>\n'
            '<h2 id="sec-1">נספח מקורות</h2>\n<table></table>\n'
        )
        out = inject_advance_footnotes_html(text, {1: "תיאור <b>מסוכן</b>."})
        assert 'class="patent-advance-note"' in out
        assert "&lt;b&gt;" in out  # escaped, never raw HTML injection
        lines = out.split("\n")
        idx = next(i for i, line in enumerate(lines) if "patent-advance-note" in line)
        assert any("נספח מקורות" in line for line in lines[idx + 1 :])


class TestInsertSectionBeforeAppendix:
    def test_md_inserts_before_appendix_heading(self):
        text = "## X\n\nbody\n\n## נספח מקורות\n\nrows\n"
        out = insert_section_before_md_appendix(text, "כותרת חדשה", "תוכן")
        lines = out.split("\n")
        assert lines.index("## כותרת חדשה") < lines.index("## נספח מקורות")

    def test_md_appends_at_end_when_no_appendix_found(self):
        text = "## X\n\nbody\n"
        out = insert_section_before_md_appendix(text, "כותרת", "תוכן")
        assert out.strip().endswith("תוכן")

    def test_html_inserts_before_appendix_heading(self):
        text = '<h2 id="sec-0">X</h2>\n<p>body</p>\n<h2 id="sec-1">נספח מקורות</h2>\n<table></table>\n'
        out = insert_section_before_html_appendix(text, "כותרת חדשה", "<svg></svg>")
        lines = out.split("\n")
        idx_new = next(i for i, line in enumerate(lines) if "כותרת חדשה" in line)
        idx_appendix = next(i for i, line in enumerate(lines) if "נספח מקורות" in line)
        assert idx_new < idx_appendix


class TestDocxTimelineShading:
    def _table_with_headers_and_rows(self, headers, rows):
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=len(headers))
        for cell, text in zip(table.rows[0].cells, headers, strict=True):
            cell.text = text
        for row_values in rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values, strict=True):
                cell.text = str(value)
        return doc, table

    def test_find_table_by_headers_matches_exact_header_row(self):
        headers = ["מספר", "סטטוס"]
        doc, table = self._table_with_headers_and_rows(headers, [["1", "פג"]])
        found = find_table_by_headers(doc, headers)
        # python-docx wraps each `w:tbl` element in a fresh Table() object on every `.tables`
        # access -- compare the underlying XML element identity, not object identity.
        assert found is not None
        assert found._tbl is table._tbl

    def test_find_table_by_headers_returns_none_when_absent(self):
        doc, _table = self._table_with_headers_and_rows(["A", "B"], [["1", "2"]])
        assert find_table_by_headers(doc, ["X", "Y"]) is None

    def test_shade_table_cell_sets_fill_in_xml(self):
        _doc, table = self._table_with_headers_and_rows(["A"], [["1"]])
        cell = table.rows[1].cells[0]
        shade_table_cell(cell, "FF0000")
        assert 'w:fill="FF0000"' in cell._tc.xml

    def test_shade_timeline_table_rows_skips_header_and_unmatched_status(self):
        headers = ["מספר", "סטטוס"]
        _doc, table = self._table_with_headers_and_rows(headers, [["1", "פג"], ["2", "בתוקף"]])
        shade_timeline_table_rows(table, status_col_index=1, status_colors={"פג": "F4CCCC"})
        header_cell = table.rows[0].cells[0]
        expired_cell = table.rows[1].cells[0]
        normal_cell = table.rows[2].cells[0]
        assert "w:fill=" not in header_cell._tc.xml
        assert 'w:fill="F4CCCC"' in expired_cell._tc.xml
        assert "w:fill=" not in normal_cell._tc.xml
