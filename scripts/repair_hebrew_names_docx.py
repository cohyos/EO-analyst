#!/usr/bin/env python
"""Round-17 repair (2026-09-17): apply :func:`eoa.report.textnorm.canonicalize_hebrew_names` to
every text run of every already-rendered ``.docx`` report under ``output/reports/versions/**`` --
the companion of ``scripts/repair_hebrew_names.py`` (DB rows) and the inline md/html fix already
applied directly to the 26 affected ``output/reports/versions/**/*.{md,html}`` files.

Deliberately NOT a re-render via the existing report builders (``eoa.report.daily.build_daily``
etc.) -- that would mean new LLM calls and fresh DB reads, i.e. a genuinely different report, not a
spelling fix to the one that was actually published. Instead this walks every already-existing text
run (body paragraphs, tables, headers/footers) in place with python-docx, setting ``run.text =
canonicalize_hebrew_names(run.text)`` wherever that changes anything -- no restructuring, no
re-analysis, cheap and safe. Re-validated with ``eoa.report.docx_builder.validate_docx`` after
every save so a corrupted write is caught immediately (and the original is restored).

Use ``--dry-run`` to see which files/runs would change without writing anything.
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

_REPO_ROOT = Path(__file__).resolve().parent.parent
_AGENT_DIR = _REPO_ROOT / "agent"
if str(_AGENT_DIR) not in sys.path:
    sys.path.insert(0, str(_AGENT_DIR))

import docx  # noqa: E402

from eoa.report.docx_builder import validate_docx  # noqa: E402
from eoa.report.textnorm import canonicalize_hebrew_names  # noqa: E402


def _patch_paragraphs(paragraphs) -> int:
    n = 0
    for p in paragraphs:
        for run in p.runs:
            if not run.text:
                continue
            new_text = canonicalize_hebrew_names(run.text)
            if new_text != run.text:
                run.text = new_text
                n += 1
    return n


def _patch_tables(tables) -> int:
    n = 0
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                n += _patch_paragraphs(cell.paragraphs)
                n += _patch_tables(cell.tables)  # nested tables, if any
    return n


def patch_docx(path: Path, *, dry_run: bool) -> int:
    doc = docx.Document(str(path))
    n = _patch_paragraphs(doc.paragraphs)
    n += _patch_tables(doc.tables)
    for section in doc.sections:
        for part in (section.header, section.footer):
            n += _patch_paragraphs(part.paragraphs)
            n += _patch_tables(part.tables)
    if n == 0 or dry_run:
        return n

    backup = path.with_suffix(path.suffix + ".bak")
    shutil.copy2(path, backup)
    try:
        doc.save(str(path))
        validate_docx(path)
    except Exception:
        shutil.copy2(backup, path)
        raise
    else:
        backup.unlink()
    return n


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--root",
        default=str(_REPO_ROOT / "output" / "reports" / "versions"),
        help="directory to scan for .docx files (default: output/reports/versions)",
    )
    args = parser.parse_args()

    root = Path(args.root)
    total_files = 0
    total_runs = 0
    for path in sorted(root.rglob("*.docx")):
        n = patch_docx(path, dry_run=args.dry_run)
        if n:
            total_files += 1
            total_runs += n
            print(f"{'would fix' if args.dry_run else 'fixed'}: {path} ({n} run(s))")

    print(f"\n{'DRY RUN: ' if args.dry_run else ''}{total_files} file(s), {total_runs} run(s) touched.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
